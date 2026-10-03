from contextlib import contextmanager
from datetime import date, datetime, time, timedelta
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import MagicMock, patch

from openpyxl import load_workbook

from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import IntegrityError, transaction
from django.test import SimpleTestCase, TestCase, override_settings
from django.urls import reverse
from django.utils import timezone

from accounts.models import CompanyProfile
from business.models import (
    CaseLockInvite,
    CaseLockProduct,
    CaseLockSubmission,
    Client,
    ClientSubscription,
    Product,
    Vendor,
    VendorPriceRequest,
    SubscriptionProductRenewal,
    SubscriptionRenewalInvite,
    SubscriptionRenewalInviteLine,
    SubscriptionRenewalInviteNewProduct,
    SubscriptionRenewalInviteRecipient,
    RenewalSLASettings,
    TeamNotification,
    annotate_internal_case_locks,
)

from business.expiry_notifications import (
    build_renewal_quote_items,
    send_renewal_confirmation_email,
    send_renewal_quotation_email,
)
from business.reminder_schedule import (
    MODE_DAILY,
    _phase,
    normalize_schedule,
    recommended_schedule,
)
from business.renewal_tokens import (
    sign_renewal_confirm_token,
    unsign_renewal_confirm_token,
)
from business.analytics import (
    empty_ticket_view,
    format_duration,
    renewal_report_charts,
    sales_dashboard_charts,
    ticket_charts_from_views,
)
from business.subscription_export import (
    NO_COMPANY_VALUE,
    build_subscription_export_bytes,
    normalize_column_keys,
)
from business.subscription_import import import_subscriptions_from_file
from business.views import (
    _renewal_invite_queryset,
    _renewal_product_groups,
)


def _product(name="Antivirus", price="1200"):
    return SimpleNamespace(name=name, unit_price=Decimal(price))


def _line(product, *, qty, name=None, company="Client Co"):
    return SimpleNamespace(
        company_name=company,
        product_name=name or product.name,
        renewal_count=qty,
        subscription=SimpleNamespace(product=product),
    )


def _new_product(product, *, qty, name=None, company="Client Co"):
    return SimpleNamespace(
        company_name=company,
        product_name=name or product.name,
        quantity=qty,
        product=product,
    )


class RenewalQuoteItemsTests(SimpleTestCase):
    def test_includes_qty_and_price_and_computes_amount(self):
        product = _product()
        items = build_renewal_quote_items([_line(product, qty=10)])
        self.assertEqual(len(items), 1)
        self.assertEqual(items[0]["product_name"], "Antivirus")
        self.assertEqual(items[0]["quantity"], 10)
        self.assertEqual(items[0]["unit_price"], Decimal("1200.00"))
        self.assertEqual(items[0]["amount"], Decimal("12000.00"))

    def test_skips_zero_quantity(self):
        items = build_renewal_quote_items([_line(_product(), qty=0)])
        self.assertEqual(items, [])

    def test_skips_zero_price(self):
        items = build_renewal_quote_items(
            [_line(_product(price="0"), qty=5)]
        )
        self.assertEqual(items, [])

    def test_skips_missing_price_product(self):
        line = SimpleNamespace(
            company_name="Co",
            product_name="Gone",
            renewal_count=2,
            subscription=SimpleNamespace(product=None),
        )
        self.assertEqual(build_renewal_quote_items([line]), [])

    def test_includes_new_products_with_price(self):
        vpn = _product("VPN", "500")
        items = build_renewal_quote_items(
            [],
            new_products=[_new_product(vpn, qty=3)],
        )
        self.assertEqual(len(items), 1)
        self.assertEqual(items[0]["amount"], Decimal("1500.00"))

    def test_mixes_renewal_and_new_and_skips_unpriced(self):
        priced = _product("Antivirus", "1200")
        unpriced = _product("Firewall", "0")
        vpn = _product("VPN", "500")
        items = build_renewal_quote_items(
            [
                _line(priced, qty=10),
                _line(_product("Backup", "800"), qty=0),
                _line(unpriced, qty=2),
            ],
            new_products=[_new_product(vpn, qty=3)],
        )
        names = [item["product_name"] for item in items]
        self.assertEqual(names, ["Antivirus", "VPN"])
        self.assertEqual(sum((item["amount"] for item in items), Decimal("0")), Decimal("13500.00"))


class SendRenewalConfirmationEmailTests(SimpleTestCase):
    def _invite(self, **overrides):
        data = {
            "pk": "invite-1",
            "company": SimpleNamespace(
                pk=1,
                company_name="Co",
                outgoing_email="sales@example.com",
                outgoing_email_app_password="secret",
                company_logo=None,
                regards_contact="",
                regards_email="",
            ),
            "client_email": "one@example.com",
            "submitted_by_email": "two@example.com",
        }
        data.update(overrides)
        return SimpleNamespace(**data)

    def test_sends_only_to_submitter_without_cc(self):
        invite = self._invite()
        with patch(
            "business.expiry_notifications.send_renewal_email"
        ) as send_mail:
            send_mail.return_value = (True, "", None)
            sent = send_renewal_confirmation_email(
                invite, lines=[], new_products=[]
            )
        self.assertTrue(sent)
        send_mail.assert_called_once()
        args, kwargs = send_mail.call_args
        self.assertEqual(args[2], ["two@example.com"])
        self.assertFalse(kwargs.get("cc_addresses"))

    def test_skips_when_submitter_missing(self):
        invite = self._invite(submitted_by_email="")
        with patch(
            "business.expiry_notifications.send_renewal_email"
        ) as send_mail:
            sent = send_renewal_confirmation_email(
                invite, lines=[], new_products=[]
            )
        self.assertFalse(sent)
        send_mail.assert_not_called()


class SendRenewalQuotationEmailTests(SimpleTestCase):
    def test_skips_when_no_priced_lines(self):
        invite = SimpleNamespace(
            pk="invite-1",
            company=SimpleNamespace(
                outgoing_email="sales@example.com",
                outgoing_email_app_password="secret",
            ),
            client_email="client@example.com",
        )
        with patch(
            "business.expiry_notifications.send_renewal_email"
        ) as send_mail:
            sent = send_renewal_quotation_email(
                invite,
                lines=[_line(_product(price="0"), qty=4)],
                new_products=[],
            )
        self.assertFalse(sent)
        send_mail.assert_not_called()


class RenewalConfirmTokenTests(SimpleTestCase):
    def test_round_trip_normalizes_email(self):
        token = sign_renewal_confirm_token("abc-id", "  Client@Example.COM ")
        parsed = unsign_renewal_confirm_token(token)
        self.assertEqual(parsed, ("abc-id", "client@example.com"))

    def test_rejects_tampered_token(self):
        token = sign_renewal_confirm_token("abc-id", "client@example.com")
        self.assertIsNone(unsign_renewal_confirm_token(token + "x"))

    def test_rejects_empty_token(self):
        self.assertIsNone(unsign_renewal_confirm_token(""))
        self.assertIsNone(unsign_renewal_confirm_token("not-a-token"))


class RenewalConfirmSignedLinkTests(TestCase):
    def setUp(self):
        user = get_user_model().objects.create_user("owner@example.com", "pass")
        company = CompanyProfile.objects.create(user=user, company_name="Co")
        self.invite = SubscriptionRenewalInvite.objects.create(
            company=company,
            client_email="one@example.com",
        )
        now = timezone.now()
        SubscriptionRenewalInviteRecipient.objects.create(
            invite=self.invite, email="one@example.com", sent_at=now
        )
        SubscriptionRenewalInviteRecipient.objects.create(
            invite=self.invite, email="two@example.com", sent_at=now
        )

    def _url(self, email=None):
        if email is None:
            return reverse(
                "renewal_confirm_public_legacy",
                kwargs={"invite_id": self.invite.pk},
            )
        token = sign_renewal_confirm_token(self.invite.pk, email)
        return reverse(
            "renewal_confirm_public",
            kwargs={"invite_id": str(self.invite.pk), "token": token},
        )

    def test_missing_token_rejected_when_recipients_exist(self):
        response = self.client.get(self._url())
        self.assertEqual(response.status_code, 403)
        self.assertContains(response, "This link is not valid", status_code=403)

    def test_valid_tokens_open_form(self):
        first = self.client.get(self._url("one@example.com"))
        self.assertEqual(first.status_code, 200)
        self.assertContains(first, "Confirm subscription renewal")
        second = self.client.get(self._url("two@example.com"))
        self.assertEqual(second.status_code, 200)
        self.assertContains(second, "Confirm subscription renewal")
        self.invite.recipients.get(email="one@example.com").refresh_from_db()
        self.assertIsNotNone(
            self.invite.recipients.get(email="one@example.com").opened_at
        )

    def test_unsent_email_token_rejected(self):
        response = self.client.get(self._url("three@example.com"))
        self.assertEqual(response.status_code, 403)

    def test_second_link_locked_after_submit(self):
        self.invite.submitted_at = timezone.now()
        self.invite.submitted_by_email = "one@example.com"
        self.invite.save(update_fields=["submitted_at", "submitted_by_email"])
        response = self.client.get(self._url("two@example.com"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Already submitted")
        self.assertContains(response, "one@example.com")


class ClientRenewalTabRowsTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user(
            "owner@example.com", "pass"
        )
        self.company = CompanyProfile.objects.create(
            user=self.user, company_name="Co"
        )
        client = Client.objects.create(
            company=self.company, name="Acme", email="acme@example.com"
        )
        existing_product = Product.objects.create(
            company=self.company, name="Antivirus"
        )
        new_product = Product.objects.create(company=self.company, name="VPN")
        subscription = ClientSubscription.objects.create(
            company=self.company,
            client=client,
            product=existing_product,
            start_date=date.today(),
            end_date=date.today() + timedelta(days=30),
            quantity=5,
        )
        now = timezone.now()
        self.invite = SubscriptionRenewalInvite.objects.create(
            company=self.company,
            client_email="acme@example.com",
            email_sent_at=now,
            submitted_at=now,
            submitted_by_email="acme@example.com",
        )
        SubscriptionRenewalInviteLine.objects.create(
            invite=self.invite,
            subscription=subscription,
            company_name="Acme",
            product_name="Antivirus",
            current_quantity=5,
            renewal_count=8,
        )
        SubscriptionRenewalInviteNewProduct.objects.create(
            invite=self.invite,
            product=new_product,
            company_name="Acme",
            product_name="VPN",
            quantity=3,
        )
        pending = SubscriptionRenewalInvite.objects.create(
            company=self.company,
            client_email="pending@example.com",
            email_sent_at=now,
        )
        SubscriptionRenewalInviteLine.objects.create(
            invite=pending,
            subscription=subscription,
            company_name="Acme",
            product_name="Antivirus",
            current_quantity=5,
        )

    def _invites(self):
        return list(_renewal_invite_queryset(self.company))

    def _rows(self):
        groups = _renewal_product_groups(self._invites())
        return [row for group in groups for row in group["rows"]]

    def test_groups_include_new_and_existing_products(self):
        rows = {row["product_name"]: row for row in self._rows()}
        self.assertEqual(rows["VPN"]["client_email"], "acme@example.com")
        self.assertTrue(rows["VPN"]["is_new"])
        self.assertEqual(rows["VPN"]["renewal_quantity"], 3)
        self.assertIsNone(rows["VPN"]["previous_quantity"])
        self.assertEqual(rows["Antivirus"]["previous_quantity"], 5)
        self.assertEqual(rows["Antivirus"]["renewal_quantity"], 8)
        self.assertFalse(rows["Antivirus"]["is_new"])
        self.assertTrue(rows["Antivirus"]["start_date"])
        self.assertTrue(rows["Antivirus"]["end_date"])
        self.assertEqual(rows["VPN"]["start_date"], "")
        self.assertEqual(rows["VPN"]["end_date"], "")
        self.assertEqual(rows["Antivirus"]["ref_id"], "")
        self.assertEqual(rows["VPN"]["ref_id"], "")

    def test_list_page_uses_single_grouped_table(self):
        self.client.force_login(self.user)
        response = self.client.get(reverse("client_renewal_list"))
        self.assertEqual(response.status_code, 200)

        groups = response.context["renewal_groups"]
        self.assertEqual(len(groups), 1)
        self.assertEqual(groups[0]["company_name"], "Acme")
        self.assertEqual(len(groups[0]["rows"]), 3)
        names = [row["product_name"] for row in groups[0]["rows"]]
        self.assertEqual(names.count("VPN"), 1)
        self.assertEqual(names.count("Antivirus"), 2)
        self.assertContains(response, "Company")
        self.assertContains(response, "Prev qty")
        self.assertContains(response, "Renewal qty")
        self.assertContains(response, "VPN")
        self.assertNotContains(response, 'class="cr-th-ref"')
        self.assertContains(response, 'id="cr-filters"')
        self.assertContains(response, "Fill form pending")
        self.assertContains(response, "Win / Lost pending")
        self.assertContains(response, "Not started")
        self.assertContains(response, "On track")
        self.assertContains(response, 'class="cr-th-sla"')
        self.assertContains(response, "SLA settings")
        self.assertContains(response, 'id="cr-sla-settings-btn"')
        self.assertContains(response, 'id="cr-sla-modal"')
        self.assertNotContains(response, 'class="cr-th-email"')
        self.assertContains(response, "Connect Zoho")
        self.assertContains(response, ">Win</a>")
        self.assertContains(response, ">Lost</button>")
        self.assertNotContains(response, "cr-deal-pill cr-status--pending")
        self.assertContains(response, "Acme")
        self.assertContains(response, 'class="cr-expand-btn"')
        self.assertContains(response, 'data-company="Acme"')
        self.assertContains(response, 'id="cr-info-modal"')
        self.assertContains(response, "Product info")
        self.assertContains(response, "Email info")
        self.assertContains(response, 'class="cr-products"')
        self.assertNotContains(response, "Show timestamps")
        self.assertNotContains(response, "cr-ts-row")
        self.assertNotContains(response, "New subscriptions")
        self.assertNotContains(response, "Renewal requests")
        self.assertNotContains(response, "All emails")
        self.assertNotContains(response, "Set status")
        self.assertNotContains(response, "Average time to close")

    def test_live_json_includes_grouped_rows(self):
        self.client.force_login(self.user)
        response = self.client.get(reverse("client_renewal_live"))
        self.assertEqual(response.status_code, 200)
        data = response.json()

        self.assertEqual(len(data["groups"]), 1)
        rows = {row["product_name"]: row for row in data["groups"][0]["rows"]}
        self.assertEqual(rows["VPN"]["renewal_quantity"], 3)
        self.assertTrue(rows["VPN"]["is_new"])
        self.assertEqual(rows["Antivirus"]["previous_quantity"], 5)
        self.assertEqual(rows["Antivirus"]["renewal_quantity"], 8)
        self.assertTrue(rows["Antivirus"]["subscription_id"])
        self.assertIn("sent", rows["Antivirus"]["timestamps"])
        self.assertTrue(rows["Antivirus"]["start_date"])
        self.assertTrue(rows["Antivirus"]["end_date"])
        self.assertEqual(rows["VPN"]["start_date"], "")
        self.assertEqual(rows["VPN"]["end_date"], "")

    def test_zoho_payload_includes_open_quote_url(self):
        self.invite.zoho_quote_url = "https://crm.zoho.com/crm/quote/1"
        self.invite.zoho_quote_id = "zquote-1"
        self.invite.zoho_quote_number = "QT-91"
        self.invite.quotation_sent_at = timezone.now()
        self.invite.save(
            update_fields=[
                "zoho_quote_url",
                "zoho_quote_id",
                "zoho_quote_number",
                "quotation_sent_at",
            ]
        )
        groups = _renewal_product_groups(
            self._invites(), zoho_connected=True, zoho_configured=True
        )
        rows = [row for group in groups for row in group["rows"]]
        antivirus = next(
            row
            for row in rows
            if row["product_name"] == "Antivirus"
            and row["client_email"] == "acme@example.com"
        )
        self.assertEqual(antivirus["zoho"]["kind"], "send")
        self.assertEqual(
            antivirus["zoho"]["quote_url"], "https://crm.zoho.com/crm/quote/1"
        )
        self.assertEqual(antivirus["ref_id"], "QT-91")
        vpn = next(row for row in rows if row["product_name"] == "VPN")
        self.assertEqual(vpn["zoho"]["quote_url"], "https://crm.zoho.com/crm/quote/1")
        self.assertEqual(vpn["ref_id"], "QT-91")
        pending = next(
            row
            for row in rows
            if row["product_name"] == "Antivirus"
            and row["client_email"] == "pending@example.com"
        )
        self.assertEqual(pending["zoho"]["kind"], "send")
        self.assertTrue(pending["zoho"]["send_url"])

    def _quote_payload(self, invite):
        from business.zoho_crm import create_quote_from_renewal_invite

        self.company.zoho_refresh_token = "refresh-token"
        self.company.save(update_fields=["zoho_refresh_token"])
        with (
            patch("business.zoho_crm.zoho_configured", return_value=True),
            patch("business.zoho_crm.find_or_create_account", return_value="acc-1"),
            patch("business.zoho_crm.find_or_create_product", return_value="prod-1"),
            patch(
                "business.zoho_crm._create_record",
                return_value={"id": "q-1", "Quoted_Items": "QT-1"},
            ) as create_record,
            patch("business.zoho_crm.requests.get") as get_quote,
        ):
            get_quote.return_value.ok = False
            create_quote_from_renewal_invite(self.company, invite)
        return create_record.call_args.args[2]

    def test_quote_uses_renewal_qty_when_present(self):
        payload = self._quote_payload(self.invite)
        self.assertEqual(payload["Product_Details"][0]["quantity"], 8)

    def test_quote_uses_current_qty_when_renewal_qty_missing(self):
        pending = SubscriptionRenewalInvite.objects.get(
            client_email="pending@example.com"
        )
        payload = self._quote_payload(pending)
        self.assertEqual(payload["Product_Details"][0]["quantity"], 5)
        self.assertIn("no renewal qty yet", payload["Description"])

    def test_sla_payload_tracks_quantity_to_quote_send(self):
        from datetime import time as dt_time

        from business.renewal_sla import sla_status_for_invite

        settings_obj = RenewalSLASettings(
            clock_mode=RenewalSLASettings.ClockMode.ALWAYS,
            target_hours=8,
            work_start=dt_time(9, 30),
            work_end=dt_time(18, 30),
        )
        start = timezone.now() - timedelta(hours=3)
        self.invite.submitted_at = start
        self.invite.quotation_sent_at = start + timedelta(hours=1)
        self.invite.zoho_quote_stage = "Draft"
        self.invite.save(
            update_fields=["submitted_at", "quotation_sent_at", "zoho_quote_stage"]
        )
        data = sla_status_for_invite(self.invite, settings_obj)
        self.assertEqual(data["zoho_status"], "draft_pending")
        self.assertEqual(data["status"], "on_track")
        self.assertAlmostEqual(data["qty_to_quote_hours"], 1.0, places=1)
        self.assertAlmostEqual(data["elapsed_hours"], 3.0, places=1)

        self.invite.zoho_quote_stage = "Delivered"
        self.invite.zoho_quote_sent_at = start + timedelta(hours=2)
        data = sla_status_for_invite(self.invite, settings_obj)
        self.assertEqual(data["zoho_status"], "sent")
        self.assertEqual(data["status"], "met")
        self.assertAlmostEqual(data["elapsed_hours"], 2.0, places=1)

    def test_save_sla_settings_from_list_page(self):
        self.client.force_login(self.user)
        response = self.client.post(
            reverse("client_renewal_list"),
            {
                "form": "renewal_sla",
                "clock_mode": "24x7",
                "target_hours": "9",
                "work_start": "09:30",
                "work_end": "18:30",
                "work_monday": "on",
                "work_tuesday": "on",
                "work_wednesday": "on",
                "work_thursday": "on",
                "work_friday": "on",
            },
        )
        self.assertEqual(response.status_code, 302)
        settings_obj = RenewalSLASettings.objects.get(company=self.company)
        self.assertEqual(settings_obj.clock_mode, "24x7")
        self.assertEqual(settings_obj.target_hours, 9)
        self.assertTrue(settings_obj.work_monday)
        self.assertFalse(settings_obj.work_saturday)

    def test_invalid_sla_settings_keep_modal_open(self):
        self.client.force_login(self.user)
        response = self.client.post(
            reverse("client_renewal_list"),
            {
                "form": "renewal_sla",
                "clock_mode": "working_hours",
                "target_hours": "8",
                "work_start": "09:30",
                "work_end": "18:30",
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.context["sla_modal_open"])
        self.assertContains(response, 'id="cr-sla-modal"')
        self.assertNotContains(response, 'id="cr-sla-modal" hidden')
        self.assertContains(response, "Tick at least one working day.")


class RenewalSLAClockTests(SimpleTestCase):
    def _working(self):
        from datetime import time as dt_time

        return RenewalSLASettings(
            clock_mode=RenewalSLASettings.ClockMode.WORKING,
            target_hours=8,
            work_start=dt_time(9, 30),
            work_end=dt_time(18, 30),
            work_monday=True,
            work_tuesday=True,
            work_wednesday=True,
            work_thursday=True,
            work_friday=True,
            work_saturday=False,
            work_sunday=False,
        )

    def test_working_hours_skip_weekend(self):
        from datetime import datetime

        from business.renewal_sla import add_clock_hours, elapsed_seconds

        tz = timezone.get_current_timezone()
        friday = timezone.make_aware(datetime(2026, 9, 18, 17, 0), tz)
        monday = timezone.make_aware(datetime(2026, 9, 21, 10, 0), tz)
        seconds = elapsed_seconds(friday, monday, self._working())
        self.assertAlmostEqual(seconds / 3600, 2.0, places=2)
        due = add_clock_hours(friday, 8, self._working())
        self.assertEqual(timezone.localtime(due).strftime("%Y-%m-%d %H:%M"), "2026-09-21 16:00")

    def test_24x7_counts_calendar_hours(self):
        from datetime import datetime, time as dt_time

        from business.renewal_sla import elapsed_seconds

        settings_obj = RenewalSLASettings(
            clock_mode=RenewalSLASettings.ClockMode.ALWAYS,
            target_hours=8,
            work_start=dt_time(9, 30),
            work_end=dt_time(18, 30),
        )
        tz = timezone.get_current_timezone()
        start = timezone.make_aware(datetime(2026, 9, 18, 17, 0), tz)
        end = timezone.make_aware(datetime(2026, 9, 21, 10, 0), tz)
        seconds = elapsed_seconds(start, end, settings_obj)
        self.assertAlmostEqual(seconds / 3600, 65.0, places=2)

    def test_quote_stage_sent_and_draft(self):
        from business.renewal_sla import quote_stage_is_sent

        self.assertFalse(quote_stage_is_sent("Draft"))
        self.assertFalse(quote_stage_is_sent("Negotiation"))
        self.assertTrue(quote_stage_is_sent("Delivered"))
        self.assertTrue(quote_stage_is_sent("Sent"))


class ReminderScheduleOverlapTests(SimpleTestCase):
    def test_recommended_schedule_has_no_overlap(self):
        normalize_schedule(recommended_schedule())

    def test_overlapping_before_expiry_windows_are_rejected(self):
        raw = [
            _phase("Early window", 90, 1, MODE_DAILY, ["09:00"]),
            _phase("Final days", 7, 1, MODE_DAILY, ["09:00"]),
        ]
        with self.assertRaises(ValidationError) as ctx:
            normalize_schedule(raw)
        msg = " ".join(ctx.exception.messages)
        self.assertIn("overlap", msg.lower())
        self.assertIn("days 7–1 left", msg)
        self.assertIn("90 to 8 days left", msg)

    def test_adjacent_before_expiry_windows_are_allowed(self):
        raw = [
            _phase("Early window", 90, 8, MODE_DAILY, ["09:00"]),
            _phase("Final days", 7, 1, MODE_DAILY, ["09:00"]),
        ]
        normalize_schedule(raw)

    def test_three_overlapping_before_windows_are_rejected(self):
        raw = [
            _phase("Early window", 90, 1, MODE_DAILY, ["09:00"]),
            _phase("Mid window", 30, 1, MODE_DAILY, ["09:00"]),
            _phase("Final days", 7, 1, MODE_DAILY, ["09:00"]),
        ]
        with self.assertRaises(ValidationError) as ctx:
            normalize_schedule(raw)
        self.assertGreaterEqual(len(ctx.exception.messages), 2)

    def test_overdue_overlap_is_rejected(self):
        raw = [
            _phase("Overdue", -1, -7, MODE_DAILY, ["09:00"]),
            _phase("Late push", -4, -10, MODE_DAILY, ["09:00"]),
        ]
        with self.assertRaises(ValidationError) as ctx:
            normalize_schedule(raw)
        msg = " ".join(ctx.exception.messages)
        self.assertIn("overlap", msg.lower())
        self.assertIn("days 4–7 late", msg)
        self.assertIn("1 to 3 days late", msg)

    def test_before_and_overdue_windows_do_not_conflict(self):
        raw = [
            _phase("Early window", 90, 1, MODE_DAILY, ["09:00"]),
            _phase("Overdue", -1, -7, MODE_DAILY, ["09:00"]),
        ]
        normalize_schedule(raw)

    def test_two_expiry_day_windows_are_rejected(self):
        raw = [
            _phase("Expiry day", 0, 0, MODE_DAILY, ["09:00"]),
            _phase("Expiry again", 0, 0, MODE_DAILY, ["09:00"]),
        ]
        with self.assertRaises(ValidationError) as ctx:
            normalize_schedule(raw)
        self.assertIn("expiry day", " ".join(ctx.exception.messages).lower())


class CaseLockFeatureTests(TestCase):
    def setUp(self):
        User = get_user_model()
        self.user = User.objects.create_user(
            email="owner@example.com", password="pass12345"
        )
        self.company = CompanyProfile.objects.create(
            user=self.user,
            company_name="Acme",
            caselock_email="caselock@example.com",
            outgoing_email="outgoing@example.com",
            outgoing_email_app_password="out-secret",
        )
        self.client.login(email="owner@example.com", password="pass12345")
        mailbox_patch = patch(
            "business.views.maybe_fetch_on_dashboard",
            return_value={"created": 0, "logs": []},
        )
        mailbox_patch.start()
        self.addCleanup(mailbox_patch.stop)

    def test_company_setup_shows_caselock_email(self):
        response = self.client.get(reverse("company_setup"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Case lock email")
        self.assertContains(response, "caselock_email")

    def test_dashboard_has_tabs(self):
        response = self.client.get(reverse("caselock_inbox_list"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Dashboard")
        self.assertContains(response, "Incoming mails")
        self.assertContains(response, "caselock@example.com")
        self.assertContains(response, "subnav-tab")

    def test_invite_and_decision_emails_use_html_layout(self):
        from business.caselock_inbox import (
            _send_invite_email,
            send_caselock_decision_email,
        )

        invite = CaseLockInvite.objects.create(
            company=self.company,
            source_message_id="<html-invite@test>",
            sender_email="buyer@example.com",
        )
        with patch(
            "business.caselock_inbox.send_outbound", return_value=(True, None)
        ) as send:
            _send_invite_email(self.company, invite, "buyer@example.com")
        self.assertEqual(send.call_args.args[0], "outgoing@example.com")
        self.assertEqual(send.call_args.args[2], ["buyer@example.com"])
        html = send.call_args.kwargs["html_body"]
        self.assertIn("SUBMIT CASE LOCK DETAILS", html)
        self.assertIn("/caselock/fill/", html)
        self.assertIn("href=", html)

        product = CaseLockProduct.objects.create(company=self.company, name="Backup")
        sub = CaseLockSubmission.objects.create(
            invite=invite,
            client_name="Buyer",
            client_gst="24AAFCN2771H2ZH",
            product=product,
            product_name="Backup",
            product_quantity=2,
            status=CaseLockSubmission.Status.APPROVED,
        )
        with patch(
            "business.caselock_inbox.send_outbound", return_value=(True, None)
        ) as send:
            self.assertTrue(send_caselock_decision_email(sub))
        html = send.call_args.kwargs["html_body"]
        self.assertIn("finalized", html.lower())
        self.assertIn("24AAFCN2771H2ZH", html)
        self.assertIn("Buyer", html)

        sub.status = CaseLockSubmission.Status.DECLINED
        sub.save(update_fields=["status"])
        with patch(
            "business.caselock_inbox.send_outbound", return_value=(True, None)
        ) as send:
            self.assertTrue(send_caselock_decision_email(sub))
        html = send.call_args.kwargs["html_body"]
        self.assertIn("declined", html.lower())
        self.assertIn("Backup", html)

    def test_incoming_mails_tab(self):
        invite = CaseLockInvite.objects.create(
            company=self.company,
            source_message_id="<one@test>",
            sender_email="buyer@example.com",
            mail_subject="Need lock",
            received_at=timezone.now(),
        )
        response = self.client.get(reverse("caselock_inbox_list") + "?tab=mails")
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "buyer@example.com")
        self.assertContains(response, "Need lock")
        self.assertContains(response, invite.sender_email)

    def test_product_add_and_public_form(self):
        add = self.client.post(
            reverse("caselock_product_add"),
            {"name": "Firewall"},
            follow=True,
        )
        self.assertEqual(add.status_code, 200)
        product = CaseLockProduct.objects.get(company=self.company, name="Firewall")
        invite = CaseLockInvite.objects.create(
            company=self.company,
            source_message_id="<two@test>",
            sender_email="buyer@example.com",
        )
        url = reverse("caselock_public_form", kwargs={"invite_id": invite.pk})
        get = self.client.get(url)
        self.assertEqual(get.status_code, 200)
        self.assertContains(get, "Search products")
        self.assertContains(get, "Firewall")
        self.assertContains(get, "cl-product-trigger")
        post = self.client.post(
            url,
            {
                "product": product.pk,
                "client_name": "Buyer Co",
                "client_gst": "22AAAAA0000A1Z5",
                "product_quantity": 3,
            },
        )
        self.assertEqual(post.status_code, 200)
        sub = CaseLockSubmission.objects.get(invite=invite)
        self.assertEqual(sub.client_name, "Buyer Co")
        self.assertEqual(sub.product_name, "Firewall")
        self.assertEqual(sub.product_quantity, 3)

        dash = self.client.get(reverse("caselock_inbox_list"))
        self.assertContains(dash, "Buyer Co")
        self.assertContains(dash, "Case lock")

    def test_public_form_requires_all_lock_fields(self):
        product = CaseLockProduct.objects.create(company=self.company, name="Laptop")
        invite = CaseLockInvite.objects.create(
            company=self.company,
            source_message_id="<req@test>",
            sender_email="buyer@example.com",
        )
        url = reverse("caselock_public_form", kwargs={"invite_id": invite.pk})
        complete = {
            "product": product.pk,
            "client_name": "Acme Corp",
            "client_gst": "22AAAAA0000A1Z5",
            "product_quantity": 10,
        }
        missing_cases = [
            {**complete, "product": ""},
            {**complete, "client_name": "  "},
            {**complete, "client_gst": ""},
            {**complete, "product_quantity": ""},
            {**complete, "product_quantity": 0},
        ]
        for payload in missing_cases:
            post = self.client.post(url, payload)
            self.assertEqual(post.status_code, 200)
            self.assertFalse(CaseLockSubmission.objects.filter(invite=invite).exists())

    def test_internal_lock_goes_to_first_submitter(self):
        product = CaseLockProduct.objects.create(company=self.company, name="Antivirus")
        first_invite = CaseLockInvite.objects.create(
            company=self.company,
            source_message_id="<a@test>",
            sender_email="first@example.com",
        )
        second_invite = CaseLockInvite.objects.create(
            company=self.company,
            source_message_id="<b@test>",
            sender_email="second@example.com",
        )
        first = CaseLockSubmission.objects.create(
            invite=first_invite,
            client_name="Acme Corp",
            client_gst="22AAAAA0000A1Z5",
            product=product,
            product_name="Antivirus",
            product_quantity=5,
        )
        second = CaseLockSubmission.objects.create(
            invite=second_invite,
            client_name="Acme Pvt Ltd",
            client_gst="22aaaaa0000a1z5",
            product=product,
            product_name="Antivirus",
            product_quantity=5,
        )
        annotate_internal_case_locks([first, second])
        self.assertTrue(first.has_internal_case_lock)
        self.assertFalse(second.has_internal_case_lock)

    def test_internal_lock_groups_by_gst_not_company_name(self):
        product = CaseLockProduct.objects.create(company=self.company, name="Firewall")
        invite_a = CaseLockInvite.objects.create(
            company=self.company,
            source_message_id="<lock-a@test>",
            sender_email="a@example.com",
        )
        invite_b = CaseLockInvite.objects.create(
            company=self.company,
            source_message_id="<lock-b@test>",
            sender_email="b@example.com",
        )
        invite_c = CaseLockInvite.objects.create(
            company=self.company,
            source_message_id="<lock-c@test>",
            sender_email="c@example.com",
        )
        same_product_qty = dict(
            product=product,
            product_name="Firewall",
            product_quantity=2,
        )
        by_company = CaseLockSubmission.objects.create(
            invite=invite_a,
            client_name="Acme Corp",
            client_gst="22AAAAA0000A1Z5",
            **same_product_qty,
        )
        other_company = CaseLockSubmission.objects.create(
            invite=invite_b,
            client_name="Beta Ltd",
            client_gst="22AAAAA0000A1Z5",
            **same_product_qty,
        )
        other_gst = CaseLockSubmission.objects.create(
            invite=invite_c,
            client_name="Acme Corp",
            client_gst="27BBBBB0000B1Z5",
            **same_product_qty,
        )
        annotate_internal_case_locks([by_company, other_company, other_gst])
        self.assertTrue(by_company.has_internal_case_lock)
        self.assertFalse(other_company.has_internal_case_lock)
        self.assertTrue(other_gst.has_internal_case_lock)

    def test_internal_lock_without_gst_uses_company_name(self):
        product = CaseLockProduct.objects.create(company=self.company, name="Router")
        invite_a = CaseLockInvite.objects.create(
            company=self.company,
            source_message_id="<nogst-a@test>",
            sender_email="a@example.com",
        )
        invite_b = CaseLockInvite.objects.create(
            company=self.company,
            source_message_id="<nogst-b@test>",
            sender_email="b@example.com",
        )
        first = CaseLockSubmission.objects.create(
            invite=invite_a,
            client_name="Acme Corp",
            client_gst="",
            product=product,
            product_name="Router",
            product_quantity=1,
        )
        second = CaseLockSubmission.objects.create(
            invite=invite_b,
            client_name="acme corp",
            client_gst="",
            product=product,
            product_name="Router",
            product_quantity=1,
        )
        annotate_internal_case_locks([first, second])
        self.assertTrue(first.has_internal_case_lock)
        self.assertFalse(second.has_internal_case_lock)

    def test_final_unblocks_same_combination_for_new_request(self):
        product = CaseLockProduct.objects.create(company=self.company, name="VPN")
        now = timezone.now()
        fields = dict(
            client_name="Acme Corp",
            client_gst="22AAAAA0000A1Z5",
            product=product,
            product_name="VPN",
            product_quantity=4,
        )
        first_invite = CaseLockInvite.objects.create(
            company=self.company,
            source_message_id="<final-a@test>",
            sender_email="first@example.com",
        )
        waiting_invite = CaseLockInvite.objects.create(
            company=self.company,
            source_message_id="<final-b@test>",
            sender_email="waiting@example.com",
        )
        new_invite = CaseLockInvite.objects.create(
            company=self.company,
            source_message_id="<final-c@test>",
            sender_email="new@example.com",
        )
        first = CaseLockSubmission.objects.create(invite=first_invite, **fields)
        waiting = CaseLockSubmission.objects.create(invite=waiting_invite, **fields)
        CaseLockSubmission.objects.filter(pk=first.pk).update(
            timestamp=now - timedelta(hours=2)
        )
        CaseLockSubmission.objects.filter(pk=waiting.pk).update(
            timestamp=now - timedelta(hours=1)
        )
        first.refresh_from_db()
        waiting.refresh_from_db()

        first.status = CaseLockSubmission.Status.APPROVED
        first.status_changed_at = now - timedelta(minutes=10)
        first.save(update_fields=["status", "status_changed_at"])

        new_sub = CaseLockSubmission.objects.create(invite=new_invite, **fields)
        CaseLockSubmission.objects.filter(pk=new_sub.pk).update(
            timestamp=now - timedelta(minutes=1)
        )
        new_sub.refresh_from_db()

        annotate_internal_case_locks([first, waiting, new_sub])
        self.assertFalse(first.has_internal_case_lock)
        self.assertFalse(waiting.has_internal_case_lock)
        self.assertTrue(new_sub.has_internal_case_lock)

    def test_decline_does_not_unblock_combination(self):
        product = CaseLockProduct.objects.create(company=self.company, name="Mail")
        now = timezone.now()
        fields = dict(
            client_name="Acme Corp",
            client_gst="22AAAAA0000A1Z5",
            product=product,
            product_name="Mail",
            product_quantity=1,
        )
        first_invite = CaseLockInvite.objects.create(
            company=self.company,
            source_message_id="<decl-a@test>",
            sender_email="first@example.com",
        )
        later_invite = CaseLockInvite.objects.create(
            company=self.company,
            source_message_id="<decl-b@test>",
            sender_email="later@example.com",
        )
        first = CaseLockSubmission.objects.create(invite=first_invite, **fields)
        later = CaseLockSubmission.objects.create(invite=later_invite, **fields)
        CaseLockSubmission.objects.filter(pk=first.pk).update(
            timestamp=now - timedelta(hours=2)
        )
        CaseLockSubmission.objects.filter(pk=later.pk).update(
            timestamp=now - timedelta(minutes=1)
        )
        first.refresh_from_db()
        later.refresh_from_db()
        first.status = CaseLockSubmission.Status.DECLINED
        first.status_changed_at = now - timedelta(minutes=30)
        first.save(update_fields=["status", "status_changed_at"])

        annotate_internal_case_locks([first, later])
        self.assertTrue(first.has_internal_case_lock)
        self.assertFalse(later.has_internal_case_lock)

    def test_final_and_decline_send_mail(self):
        product = CaseLockProduct.objects.create(company=self.company, name="Backup")
        invite = CaseLockInvite.objects.create(
            company=self.company,
            source_message_id="<c@test>",
            sender_email="buyer@example.com",
        )
        sub = CaseLockSubmission.objects.create(
            invite=invite,
            client_name="Buyer",
            product=product,
            product_name="Backup",
            product_quantity=1,
        )
        with patch("business.views.send_caselock_decision_email", return_value=True) as send:
            response = self.client.post(
                reverse("caselock_submission_status_update", args=[sub.pk]),
                {"action": "final"},
            )
            self.assertEqual(response.status_code, 302)
            send.assert_called_once()
        sub.refresh_from_db()
        self.assertEqual(sub.status, CaseLockSubmission.Status.APPROVED)
        self.assertIsNotNone(sub.status_changed_at)

        invite2 = CaseLockInvite.objects.create(
            company=self.company,
            source_message_id="<d@test>",
            sender_email="buyer2@example.com",
        )
        sub2 = CaseLockSubmission.objects.create(
            invite=invite2,
            client_name="Other",
            product=product,
            product_name="Backup",
            product_quantity=2,
        )
        with patch("business.views.send_caselock_decision_email", return_value=True) as send:
            response = self.client.post(
                reverse("caselock_submission_status_update", args=[sub2.pk]),
                {"action": "decline"},
            )
            self.assertEqual(response.status_code, 302)
            send.assert_called_once()
        sub2.refresh_from_db()
        self.assertEqual(sub2.status, CaseLockSubmission.Status.DECLINED)

    def test_live_json_returns_incoming_mails(self):
        invite = CaseLockInvite.objects.create(
            company=self.company,
            source_message_id="<live@test>",
            sender_email="buyer@example.com",
            mail_subject="Need lock",
            received_at=timezone.now(),
        )
        response = self.client.get(reverse("caselock_inbox_live"))
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertTrue(data["ok"])
        self.assertEqual(len(data["mail_rows"]), 1)
        self.assertEqual(data["mail_rows"][0]["from"], "buyer@example.com")
        self.assertEqual(data["mail_rows"][0]["subject"], "Need lock")
        self.assertEqual(data["mail_rows"][0]["id"], str(invite.pk))
        self.assertFalse(data["mail_rows"][0]["form_submitted"])

    def test_incoming_mails_page_has_live_script(self):
        response = self.client.get(reverse("caselock_inbox_list") + "?tab=mails")
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, reverse("caselock_inbox_live"))
        self.assertContains(response, "cl-mails-tbody")
        self.assertContains(response, "cl-live-meta")

    @override_settings(DEBUG=False)
    def test_production_live_does_not_poll_mailbox(self):
        with patch("business.views.maybe_fetch_on_dashboard") as fetch:
            response = self.client.get(reverse("caselock_inbox_live"))
        self.assertEqual(response.status_code, 200)
        fetch.assert_not_called()
        self.assertFalse(response.json()["mailbox_poll_on_request"])

    @override_settings(DEBUG=True)
    def test_local_live_polls_mailbox(self):
        with patch(
            "business.views.maybe_fetch_on_dashboard",
            return_value={"created": 0, "logs": []},
        ) as fetch:
            response = self.client.get(reverse("caselock_inbox_live"))
        self.assertEqual(response.status_code, 200)
        fetch.assert_called_once()
        self.assertTrue(response.json()["mailbox_poll_on_request"])

    def test_skips_exchange_bounce_inbound(self):
        from business.caselock_inbox import (
            _create_invite_and_notify,
            _should_skip_inbound,
        )

        bounce = (
            "MicrosoftExchange329e71ec88ae4615bbc36ab6ce41109e@natrajinfo.in"
        )
        self.assertTrue(
            _should_skip_inbound(
                self.company, bounce, "Undeliverable: Case lock details"
            )
        )
        self.assertTrue(
            _should_skip_inbound(
                self.company, "caselock@example.com", "Need lock"
            )
        )
        self.assertTrue(
            _should_skip_inbound(
                self.company, "outgoing@example.com", "Need lock"
            )
        )
        self.assertFalse(
            _should_skip_inbound(self.company, "buyer@example.com", "Need lock")
        )
        created = _create_invite_and_notify(
            self.company,
            "<bounce@test>",
            bounce,
            "Undeliverable: Case lock details",
        )
        self.assertFalse(created)
        self.assertFalse(
            CaseLockInvite.objects.filter(source_message_id="<bounce@test>").exists()
        )

    def test_invite_send_uses_outgoing_not_caselock_mailbox(self):
        from business.caselock_inbox import _create_invite_and_notify

        with patch(
            "business.caselock_inbox.send_outbound", return_value=(True, None)
        ) as send:
            created = _create_invite_and_notify(
                self.company,
                "<real@test>",
                "buyer@example.com",
                "Need lock",
            )
        self.assertTrue(created)
        self.assertEqual(send.call_args.args[0], "outgoing@example.com")
        self.assertNotEqual(send.call_args.args[0], "caselock@example.com")
        invite = CaseLockInvite.objects.get(source_message_id="<real@test>")
        self.assertIsNotNone(invite.email_sent_at)

    def test_no_outbound_without_outgoing_email(self):
        from business.caselock_inbox import _send_mail_via_company

        self.company.outgoing_email = ""
        self.company.save(update_fields=["outgoing_email"])
        with patch("business.caselock_inbox.send_outbound") as send:
            ok = _send_mail_via_company(
                self.company,
                "buyer@example.com",
                "Case lock details",
                "body",
                "<p>html</p>",
            )
        self.assertFalse(ok)
        send.assert_not_called()


class SubscriptionRenewVendorTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user(
            "owner@example.com", "pass"
        )
        self.company = CompanyProfile.objects.create(
            user=self.user, company_name="Co"
        )
        self.client_obj = Client.objects.create(
            company=self.company, name="OMSGuru"
        )
        self.product = Product.objects.create(
            company=self.company, name="Microsoft 365"
        )
        self.vendor = Vendor.objects.create(
            company=self.company, name="Redington"
        )
        self.other_vendor = Vendor.objects.create(
            company=self.company, name="Ingram"
        )
        self.subscription = ClientSubscription.objects.create(
            company=self.company,
            client=self.client_obj,
            product=self.product,
            vendor=self.vendor,
            start_date=date(2025, 9, 4),
            end_date=date(2026, 9, 3),
            quantity=2,
            billing_amount=Decimal("3137.29"),
        )
        self.client.force_login(self.user)

    def _post(self, vendor):
        return self.client.post(
            reverse("subscription_renew", kwargs={"pk": self.subscription.pk}),
            {
                "start_date": "2026-09-04",
                "end_date": "2027-09-03",
                "quantity": "2",
                "billing_amount": "3137.29",
                "vendor": vendor,
            },
        )

    def test_form_lists_existing_vendors(self):
        response = self.client.get(
            reverse("subscription_renew", kwargs={"pk": self.subscription.pk})
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "New vendor")
        self.assertContains(response, "Redington")
        self.assertContains(response, "Ingram")
        self.assertContains(response, "+ Add new vendor")
        self.assertNotContains(response, "New vendor name")
        self.assertContains(
            response,
            f'<option value="{self.vendor.pk}" selected>',
            html=False,
        )

    def test_select_existing_vendor(self):
        response = self._post(str(self.other_vendor.pk))
        self.assertEqual(response.status_code, 302)
        self.subscription.refresh_from_db()
        self.assertEqual(self.subscription.vendor_id, self.other_vendor.pk)
        self.assertEqual(self.subscription.deal_status, ClientSubscription.DealStatus.WIN)
        self.assertTrue(self.subscription.win_form_completed)
        self.assertEqual(self.subscription.renewal_draft_vendor_id, self.other_vendor.pk)
        self.assertEqual(self.subscription.renewal_draft_quantity, 2)

    def test_add_new_vendor_redirects_to_add_page(self):
        response = self._post("__new__")
        self.assertEqual(response.status_code, 302)
        self.assertIn(reverse("vendor_add"), response["Location"])
        self.subscription.refresh_from_db()
        self.assertEqual(self.subscription.vendor_id, self.vendor.pk)
        self.assertFalse(
            Vendor.objects.filter(
                company=self.company, name="ABC Distributors"
            ).exists()
        )

    def test_add_vendor_returns_to_renew_with_new_vendor(self):
        renew_url = reverse(
            "subscription_renew", kwargs={"pk": self.subscription.pk}
        )
        response = self.client.post(
            reverse("vendor_add"),
            {"name": "ABC Distributors", "next": renew_url},
        )
        self.assertEqual(response.status_code, 302)
        created = Vendor.objects.get(
            company=self.company, name="ABC Distributors"
        )
        self.assertIn(renew_url, response["Location"])
        self.assertIn(f"preselect_vendor={created.pk}", response["Location"])
        renew = self.client.get(response["Location"])
        self.assertEqual(renew.status_code, 200)
        self.assertContains(
            renew,
            f'<option value="{created.pk}" selected>',
            html=False,
        )


class ClientRenewalDealStatusTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user(
            "owner@example.com", "pass"
        )
        self.company = CompanyProfile.objects.create(
            user=self.user, company_name="Co"
        )
        self.client_obj = Client.objects.create(
            company=self.company, name="Hyfun", email="hyfun@example.com"
        )
        self.product = Product.objects.create(
            company=self.company, name="Microsoft 365 Business Basic"
        )
        self.vendor = Vendor.objects.create(
            company=self.company, name="Redington"
        )
        self.subscription = ClientSubscription.objects.create(
            company=self.company,
            client=self.client_obj,
            product=self.product,
            vendor=self.vendor,
            start_date=date(2024, 9, 11),
            end_date=date(2026, 9, 10),
            quantity=1,
            billing_amount=Decimal("7529.80"),
        )
        self.invite = SubscriptionRenewalInvite.objects.create(
            company=self.company,
            client_email="hyfun@example.com",
            email_sent_at=timezone.now(),
        )
        SubscriptionRenewalInviteLine.objects.create(
            invite=self.invite,
            subscription=self.subscription,
            company_name="Hyfun Agrilink Private Limited",
            product_name=self.product.name,
            current_quantity=1,
        )
        self.client.force_login(self.user)
        self.url = reverse(
            "client_renewal_set_deal_status", kwargs={"pk": self.invite.pk}
        )

    def test_win_click_opens_form_without_marking_win(self):
        response = self.client.post(
            self.url,
            {
                "deal_status": "win",
                "subscription_id": self.subscription.pk,
            },
        )
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertTrue(data["ok"])
        self.assertEqual(data["open_win_form"], self.subscription.pk)
        self.assertFalse(data["deal_form_empty"])
        item = data["items"][0]
        self.assertEqual(item["deal_status"], "")
        self.assertFalse(item["form_empty"])
        self.subscription.refresh_from_db()
        self.assertEqual(self.subscription.deal_status, "")
        self.assertFalse(self.subscription.win_form_completed)
        self.assertEqual(self.subscription.quantity, 1)
        self.assertEqual(self.subscription.start_date, date(2024, 9, 11))

    def test_unsaved_forms_use_client_quantity_and_one_year(self):
        self.invite.submitted_at = timezone.now()
        self.invite.save(update_fields=["submitted_at"])
        line = self.invite.lines.get()
        line.renewal_count = 4
        line.save(update_fields=["renewal_count"])

        response = self.client.post(
            self.url,
            {
                "deal_status": "win",
                "subscription_id": self.subscription.pk,
            },
        )
        draft = response.json()["items"][0]["draft"]
        self.assertEqual(draft["quantity"], 4)
        self.assertEqual(draft["start_date"], "2026-09-11")
        self.assertEqual(draft["end_date"], "2027-09-11")

        renew = self.client.get(
            reverse("subscription_renew", kwargs={"pk": self.subscription.pk})
        )
        self.assertEqual(renew.status_code, 200)
        form = renew.context["form"]
        self.assertEqual(form.initial["quantity"], 4)
        self.assertEqual(form.initial["start_date"], date(2026, 9, 11))
        self.assertEqual(form.initial["end_date"], date(2027, 9, 11))

    def test_cancel_is_rejected(self):
        response = self.client.post(
            self.url,
            {
                "deal_status": "cancel",
                "subscription_id": self.subscription.pk,
            },
        )
        self.assertEqual(response.status_code, 400)
        self.subscription.refresh_from_db()
        self.assertEqual(self.subscription.deal_status, "")

    def test_win_form_applies_yearly_renewal(self):
        self.client.post(
            self.url,
            {
                "deal_status": "win",
                "subscription_id": self.subscription.pk,
            },
        )
        response = self.client.post(
            self.url,
            {
                "action": "save_win_form",
                "subscription_id": self.subscription.pk,
                "start_date": "2026-09-11",
                "end_date": "2028-09-09",
                "quantity": "3",
                "billing_amount": "22589.40",
                "vendor": str(self.vendor.pk),
                "invoice_number": "INV-778",
                "invoice_pdf": SimpleUploadedFile(
                    "renewal.pdf",
                    b"%PDF-1.4 test",
                    content_type="application/pdf",
                ),
            },
        )
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertFalse(data["deal_form_empty"])
        self.assertFalse(data["items"][0]["form_empty"])
        self.assertEqual(data["items"][0]["deal_status"], "win")
        self.assertEqual(data["items"][0]["draft"]["invoice_number"], "INV-778")
        self.subscription.refresh_from_db()
        self.assertEqual(self.subscription.deal_status, "win")
        self.assertTrue(self.subscription.win_form_completed)
        self.assertEqual(self.subscription.start_date, date(2026, 9, 11))
        self.assertEqual(self.subscription.end_date, date(2027, 9, 11))
        self.assertEqual(self.subscription.quantity, 3)
        self.assertEqual(self.subscription.billing_amount, Decimal("22589.40"))
        self.assertEqual(self.subscription.invoice_number, "INV-778")
        self.assertEqual(self.subscription.renewal_draft_start_date, date(2026, 9, 11))
        self.assertEqual(self.subscription.renewal_draft_end_date, date(2027, 9, 11))
        self.assertEqual(self.subscription.renewal_draft_quantity, 3)
        self.assertEqual(
            self.subscription.renewal_draft_billing_amount, Decimal("22589.40")
        )
        self.assertEqual(self.subscription.renewal_draft_invoice_number, "INV-778")
        self.assertTrue(self.subscription.renewal_draft_invoice_pdf)

        renew = self.client.get(
            reverse("subscription_renew", kwargs={"pk": self.subscription.pk})
        )
        self.assertEqual(renew.status_code, 200)
        form = renew.context["form"]
        self.assertEqual(form.initial["start_date"], date(2026, 9, 11))
        self.assertEqual(form.initial["end_date"], date(2027, 9, 11))
        self.assertEqual(form.initial["quantity"], 3)
        self.assertEqual(form.initial["billing_amount"], Decimal("22589.40"))
        self.assertEqual(form.initial["vendor"], str(self.vendor.pk))
        self.assertEqual(form.initial["invoice_number"], "INV-778")
        self.assertContains(renew, "New invoice number")
        self.assertContains(renew, "New invoice PDF")
        self.assertContains(renew, "New vendor")

    def test_win_form_duration_of_two_years(self):
        response = self.client.post(
            self.url,
            {
                "action": "save_win_form",
                "subscription_id": self.subscription.pk,
                "start_date": "2026-08-30",
                "end_date": "2027-08-30",
                "duration": "2",
                "quantity": "1",
                "billing_amount": "100.00",
                "vendor": str(self.vendor.pk),
            },
        )
        self.assertEqual(response.status_code, 200)
        self.subscription.refresh_from_db()
        self.assertEqual(self.subscription.start_date, date(2026, 8, 30))
        self.assertEqual(self.subscription.end_date, date(2028, 8, 30))

    def test_win_form_can_draft_other_plan_without_changing_live_product(self):
        cheap = Product.objects.create(
            company=self.company,
            name="Microsoft 365 Business Standard",
            unit_price=Decimal("100.00"),
        )
        premium = Product.objects.create(
            company=self.company,
            name="Microsoft 365 E5",
            unit_price=Decimal("5000.00"),
        )
        self.product.unit_price = Decimal("50.00")
        self.product.save(update_fields=["unit_price"])
        self.client.post(
            self.url,
            {
                "deal_status": "win",
                "subscription_id": self.subscription.pk,
            },
        )
        response = self.client.post(
            self.url,
            {
                "action": "save_win_form",
                "subscription_id": self.subscription.pk,
                "start_date": "2026-09-11",
                "end_date": "2027-09-10",
                "quantity": "2",
                "billing_amount": "9000.00",
                "vendor": str(self.vendor.pk),
                "renew_other_plan": "1",
                "product": str(cheap.pk),
            },
        )
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(
            [row["name"] for row in data["products"]],
            [
                self.product.name,
                cheap.name,
                premium.name,
            ],
        )
        draft = data["items"][0]["draft"]
        self.assertTrue(draft["renew_other_plan"])
        self.assertEqual(draft["product"], str(cheap.pk))
        self.assertEqual(draft["product_name"], cheap.name)
        self.subscription.refresh_from_db()
        self.assertEqual(self.subscription.product_id, cheap.pk)
        self.assertEqual(self.subscription.deal_status, "win")
        self.assertEqual(self.subscription.renewal_draft_product_id, cheap.pk)
        self.assertEqual(self.subscription.quantity, 2)
        self.assertEqual(self.subscription.start_date, date(2026, 9, 11))
        self.assertEqual(self.subscription.end_date, date(2027, 9, 11))
        change = SubscriptionProductRenewal.objects.get(
            subscription=self.subscription
        )
        self.assertEqual(change.old_product_name, self.product.name)
        self.assertEqual(change.new_product_name, cheap.name)
        self.assertTrue(change.applied)

        list_page = self.client.get(reverse("client_renewal_list"))
        self.assertContains(list_page, "New product renewal")
        self.assertContains(list_page, "sla-switch")
        self.assertContains(list_page, "cr-win-product-cols")
        history = self.client.get(reverse("client_renewal_product_changes"))
        self.assertEqual(history.status_code, 200)
        history_data = history.json()
        self.assertTrue(history_data["ok"])
        self.assertEqual(len(history_data["items"]), 1)
        self.assertEqual(history_data["items"][0]["old_product_name"], self.product.name)
        self.assertEqual(history_data["items"][0]["new_product_name"], cheap.name)
        self.assertTrue(history_data["items"][0]["applied"])

        renew = self.client.get(
            reverse("subscription_renew", kwargs={"pk": self.subscription.pk})
        )
        self.assertEqual(renew.status_code, 200)
        self.assertContains(renew, "New renewal product")
        self.assertContains(renew, cheap.name)

        saved = self.client.post(
            reverse("subscription_renew", kwargs={"pk": self.subscription.pk}),
            {
                "start_date": "2026-09-11",
                "end_date": "2027-09-10",
                "quantity": "2",
                "billing_amount": "9000.00",
                "vendor": str(self.vendor.pk),
                "product": str(cheap.pk),
                "renew_other_plan": "1",
            },
        )
        self.assertEqual(saved.status_code, 302)
        self.subscription.refresh_from_db()
        self.assertEqual(self.subscription.product_id, cheap.pk)
        self.assertEqual(self.subscription.product.name, cheap.name)
        change.refresh_from_db()
        self.assertTrue(change.applied)
        self.assertEqual(change.old_product_name, self.product.name)
        self.assertEqual(change.new_product_name, cheap.name)
        self.assertEqual(SubscriptionProductRenewal.objects.count(), 1)

    def test_same_plan_renewal_does_not_create_product_change(self):
        self.client.post(
            self.url,
            {
                "deal_status": "win",
                "subscription_id": self.subscription.pk,
            },
        )
        response = self.client.post(
            self.url,
            {
                "action": "save_win_form",
                "subscription_id": self.subscription.pk,
                "start_date": "2026-09-11",
                "end_date": "2027-09-10",
                "quantity": "2",
                "billing_amount": "9000.00",
                "vendor": str(self.vendor.pk),
                "renew_other_plan": "0",
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(SubscriptionProductRenewal.objects.count(), 0)

    def test_add_subscription_form_has_optional_invoice_fields(self):
        response = self.client.get(reverse("subscription_add"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Invoice number")
        self.assertContains(response, "Invoice PDF copy")
        self.assertContains(response, 'enctype="multipart/form-data"')

    def test_renew_save_sets_win_and_clears_form_empty(self):
        self.subscription.deal_status = ClientSubscription.DealStatus.WIN
        self.subscription.win_form_completed = False
        self.subscription.save(update_fields=["deal_status", "win_form_completed"])
        response = self.client.post(
            reverse("subscription_renew", kwargs={"pk": self.subscription.pk}),
            {
                "start_date": "2026-09-11",
                "end_date": "2028-09-09",
                "quantity": "3",
                "billing_amount": "22589.40",
                "vendor": str(self.vendor.pk),
                "invoice_number": "INV-900",
            },
        )
        self.assertEqual(response.status_code, 302)
        self.subscription.refresh_from_db()
        self.assertEqual(self.subscription.deal_status, "win")
        self.assertTrue(self.subscription.win_form_completed)
        self.assertEqual(self.subscription.quantity, 3)
        self.assertEqual(self.subscription.start_date, date(2026, 9, 11))
        self.assertEqual(self.subscription.renewal_draft_quantity, 3)
        self.assertEqual(self.subscription.invoice_number, "INV-900")
        state = self.invite.deal_table_state()
        self.assertFalse(state["form_empty"])
        self.assertFalse(state["pending"])

    def test_lost_remark_is_optional_and_schedules_next_year_reminders(self):
        response = self.client.post(
            self.url,
            {
                "deal_status": "lost",
                "subscription_id": self.subscription.pk,
                "lost_remark": "  Price too high  ",
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["items"][0]["deal_status"], "lost")
        self.subscription.refresh_from_db()
        self.assertEqual(self.subscription.deal_status, "lost")
        self.assertEqual(self.subscription.lost_remark, "Price too high")
        self.assertEqual(self.subscription.lost_reminder_resume_on, date(2027, 8, 26))
        self.assertIsNone(self.subscription.reminder_target_end(date(2026, 9, 24)))
        self.assertEqual(
            self.subscription.reminder_target_end(date(2027, 8, 26)),
            date(2027, 9, 10),
        )
        self.assertIsNone(self.subscription.reminder_target_end(date(2027, 8, 25)))

    def test_lost_without_remark_still_marks_lost(self):
        response = self.client.post(
            self.url,
            {
                "deal_status": "lost",
                "subscription_id": self.subscription.pk,
            },
        )
        self.assertEqual(response.status_code, 200)
        self.subscription.refresh_from_db()
        self.assertEqual(self.subscription.deal_status, "lost")
        self.assertEqual(self.subscription.lost_remark, "")
        self.assertEqual(self.subscription.lost_reminder_resume_on, date(2027, 8, 26))

    def test_october_end_resumes_reminders_fifteen_days_before_next_year(self):
        self.subscription.end_date = date(2026, 10, 30)
        self.subscription.save(update_fields=["end_date"])
        self.client.post(
            self.url,
            {
                "deal_status": "lost",
                "subscription_id": self.subscription.pk,
            },
        )
        self.subscription.refresh_from_db()
        self.assertEqual(self.subscription.lost_reminder_resume_on, date(2027, 10, 15))
        self.assertIsNone(self.subscription.reminder_target_end(date(2027, 10, 14)))
        self.assertEqual(
            self.subscription.reminder_target_end(date(2027, 10, 15)),
            date(2027, 10, 30),
        )

    def test_lost_subscription_returns_to_reminder_query_on_resume_date(self):
        from business.expiry_notifications import _active_expiring_for_client_reminders

        self.subscription.end_date = date(2026, 10, 30)
        self.subscription.save(update_fields=["end_date"])
        self.client.post(
            self.url,
            {
                "deal_status": "lost",
                "subscription_id": self.subscription.pk,
            },
        )
        before = _active_expiring_for_client_reminders(
            self.company, 90, date(2026, 10, 1), min_days=-31
        )
        during = _active_expiring_for_client_reminders(
            self.company, 90, date(2027, 10, 15), min_days=-31
        )
        self.assertNotIn(self.subscription.pk, [sub.pk for sub in before])
        self.assertIn(self.subscription.pk, [sub.pk for sub in during])

    def test_required_renew_page_removes_row_after_submit(self):
        renew_url = reverse("subscription_renew", kwargs={"pk": self.subscription.pk})
        listing = self.client.get(reverse("client_renewal_list"))
        self.assertContains(listing, renew_url)
        self.assertContains(listing, "required=1")

        form_page = self.client.get(
            renew_url,
            {"required": "1", "next": reverse("client_renewal_list")},
        )
        self.assertEqual(form_page.status_code, 200)
        self.assertContains(form_page, "Renew other plan")
        self.assertContains(form_page, "Duration")
        self.assertContains(form_page, "years")
        self.assertContains(form_page, "Renewal quantity")
        self.assertNotContains(form_page, ">Cancel</a>")

        saved = self.client.post(
            renew_url,
            {
                "required": "1",
                "next": reverse("client_renewal_list"),
                "start_date": "2026-10-07",
                "duration": "2",
                "quantity": "2",
                "billing_amount": "10.00",
                "vendor": str(self.vendor.pk),
            },
        )
        self.assertEqual(saved.status_code, 302)
        self.subscription.refresh_from_db()
        self.assertEqual(self.subscription.deal_status, "win")
        self.assertTrue(self.subscription.win_form_completed)
        self.assertEqual(self.subscription.start_date, date(2026, 10, 7))
        self.assertEqual(self.subscription.end_date, date(2028, 10, 7))
        again = self.client.get(reverse("client_renewal_list"))
        self.assertContains(again, self.product.name)
        self.assertNotContains(again, f"/subscriptions/{self.subscription.pk}/renew/")

    def test_renewal_list_uses_same_choice_colour_and_lost_remark_modal(self):
        page = self.client.get(reverse("client_renewal_list"))
        self.assertEqual(page.status_code, 200)
        self.assertContains(page, 'class="cr-th-deal"')
        self.assertContains(page, 'id="cr-lost-modal"')
        self.assertContains(page, "Add remark")
        self.assertContains(page, 'name="duration"')
        self.assertContains(page, "years")
        self.assertNotContains(page, ">Yearly<")
        self.assertContains(page, "Renewal quantity")
        self.assertNotContains(page, "cr-win-back")


class DashboardChartTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user(
            "owner@example.com", "pass"
        )
        self.company = CompanyProfile.objects.create(
            user=self.user, company_name="Co"
        )
        self.client_obj = Client.objects.create(
            company=self.company, name="Hyfun", email="hyfun@example.com"
        )
        self.product = Product.objects.create(
            company=self.company, name="Microsoft 365 Business Basic"
        )
        self.vendor = Vendor.objects.create(company=self.company, name="Redington")
        today = timezone.localdate()
        self.win_sub = ClientSubscription.objects.create(
            company=self.company,
            client=self.client_obj,
            product=self.product,
            vendor=self.vendor,
            start_date=today - timedelta(days=365),
            end_date=today + timedelta(days=5),
            quantity=10,
            billing_amount=Decimal("1000.00"),
            deal_status=ClientSubscription.DealStatus.WIN,
        )
        self.lost_sub = ClientSubscription.objects.create(
            company=self.company,
            client=self.client_obj,
            product=self.product,
            vendor=self.vendor,
            start_date=today - timedelta(days=365),
            end_date=today + timedelta(days=8),
            quantity=2,
            billing_amount=Decimal("200.00"),
            deal_status=ClientSubscription.DealStatus.LOST,
        )
        self.pending_sub = ClientSubscription.objects.create(
            company=self.company,
            client=self.client_obj,
            product=self.product,
            vendor=self.vendor,
            start_date=today - timedelta(days=365),
            end_date=today + timedelta(days=3),
            quantity=4,
            billing_amount=Decimal("400.00"),
        )
        invite = SubscriptionRenewalInvite.objects.create(
            company=self.company,
            client_email="hyfun@example.com",
            email_sent_at=timezone.now(),
            submitted_at=timezone.now(),
        )
        SubscriptionRenewalInviteLine.objects.create(
            invite=invite,
            subscription=self.pending_sub,
            company_name="Hyfun",
            product_name=self.product.name,
            current_quantity=4,
        )
        SubscriptionRenewalInviteNewProduct.objects.create(
            invite=invite,
            product=self.product,
            company_name="Hyfun",
            product_name="VPN",
            quantity=7,
        )
        self.client.force_login(self.user)

    def test_sales_charts_count_win_lost_pending_and_new_qty(self):
        charts = sales_dashboard_charts(self.company, today=timezone.localdate())
        self.assertEqual(charts["win"], 1)
        self.assertEqual(charts["lost"], 1)
        self.assertEqual(charts["pending_deals"], 1)
        self.assertEqual(charts["win_rate"], 50)
        self.assertGreaterEqual(charts["new_qty_month"], 7)
        self.assertEqual(len(charts["pipeline_labels"]), 14)
        self.assertEqual(len(charts["qty_month_labels"]), 6)

    def test_dashboard_renders_chart_markup(self):
        response = self.client.get(reverse("dashboard"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Deal outcome")
        self.assertContains(response, "Pending renewals")
        self.assertContains(response, "New subscription qty")
        self.assertContains(response, "chart-deal-outcome")
        self.assertContains(response, "sales-charts-data")

    @patch(
        "business.helpdesk_inbox.maybe_fetch_on_dashboard",
        return_value={"created": 0, "logs": []},
    )
    def test_support_tickets_renders_charts(self, _fetch):
        response = self.client.get(reverse("support_ticket_list"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Avg resolution time")
        self.assertContains(response, "SLA breached")
        self.assertContains(response, "chart-status-mix")
        self.assertContains(response, "ticket-charts-data")


class RenewalReportTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user(
            "renewal-report@example.com", "pass"
        )
        self.company = CompanyProfile.objects.create(
            user=self.user, company_name="Report Co"
        )
        self.client_obj = Client.objects.create(
            company=self.company, name="Acme Steel", email="acme@example.com"
        )
        self.product = Product.objects.create(company=self.company, name="Firewall")
        self.add_on = Product.objects.create(company=self.company, name="Backup Cloud")
        self.vendor = Vendor.objects.create(company=self.company, name="Vendor")
        today = timezone.localdate()
        tz = timezone.get_current_timezone()
        eligible = timezone.make_aware(
            datetime.combine(today + timedelta(days=10) - timedelta(days=30), time.min),
            tz,
        )
        sent = eligible + timedelta(days=2)
        deal_at = sent + timedelta(days=4)

        self.pending_sub = ClientSubscription.objects.create(
            company=self.company,
            client=self.client_obj,
            product=self.product,
            vendor=self.vendor,
            start_date=today - timedelta(days=355),
            end_date=today + timedelta(days=10),
            quantity=4,
            billing_amount=Decimal("4000.00"),
        )
        self.silent_sub = ClientSubscription.objects.create(
            company=self.company,
            client=self.client_obj,
            product=self.product,
            vendor=self.vendor,
            start_date=today - timedelta(days=400),
            end_date=today - timedelta(days=12),
            quantity=2,
            billing_amount=Decimal("800.00"),
        )
        self.win_sub = ClientSubscription.objects.create(
            company=self.company,
            client=self.client_obj,
            product=self.product,
            vendor=self.vendor,
            start_date=today - timedelta(days=365),
            end_date=today + timedelta(days=10),
            quantity=3,
            billing_amount=Decimal("1500.00"),
            deal_status=ClientSubscription.DealStatus.WIN,
            deal_status_set_at=deal_at,
        )
        self.lost_sub = ClientSubscription.objects.create(
            company=self.company,
            client=self.client_obj,
            product=self.product,
            vendor=self.vendor,
            start_date=today - timedelta(days=365),
            end_date=today - timedelta(days=3),
            quantity=1,
            billing_amount=Decimal("250.00"),
            deal_status=ClientSubscription.DealStatus.LOST,
            deal_status_set_at=timezone.now(),
        )
        pending_invite = SubscriptionRenewalInvite.objects.create(
            company=self.company,
            client_email="acme@example.com",
            email_sent_at=sent,
            read_at=sent + timedelta(hours=6),
            submitted_at=sent + timedelta(days=1),
        )
        SubscriptionRenewalInviteLine.objects.create(
            invite=pending_invite,
            subscription=self.pending_sub,
            company_name="Acme Steel",
            product_name=self.product.name,
            current_quantity=4,
            renewal_count=6,
        )
        SubscriptionRenewalInviteNewProduct.objects.create(
            invite=pending_invite,
            product=self.add_on,
            company_name="Acme Steel",
            product_name="Backup Cloud",
            quantity=5,
        )
        win_invite = SubscriptionRenewalInvite.objects.create(
            company=self.company,
            client_email="acme@example.com",
            email_sent_at=sent,
            submitted_at=sent + timedelta(days=1),
        )
        SubscriptionRenewalInviteLine.objects.create(
            invite=win_invite,
            subscription=self.win_sub,
            company_name="Acme Steel",
            product_name=self.product.name,
            current_quantity=3,
            renewal_count=3,
        )
        self.client.force_login(self.user)

    def test_report_counts_pending_silent_lost_and_new_products(self):
        report = renewal_report_charts(self.company, range_key="90")
        self.assertEqual(report["pending_renewals"], 1)
        self.assertEqual(report["silent_expired"], 1)
        self.assertEqual(report["lost_count"], 1)
        self.assertEqual(report["win_count"], 1)
        self.assertEqual(report["resub_rate"], 50)
        self.assertEqual(report["did_not_renew"], 2)
        self.assertEqual(report["new_product_clients"], 1)
        self.assertEqual(report["new_product_qty"], 5)
        self.assertEqual(report["qty_up"], 1)
        self.assertEqual(report["qty_same"], 1)
        self.assertEqual(report["avg_send_delay"], "2d")
        self.assertEqual(report["avg_deal_delay"], "4d")
        self.assertEqual(report["charts"]["funnel_values"][0], 2)
        remaining_clients = [row["client"] for row in report["remaining_rows"]]
        self.assertIn("Acme Steel", remaining_clients)
        self.assertTrue(report["new_product_rows"])
        reasons = {row["reason"] for row in report["did_not_renew_rows"]}
        self.assertIn("Lost", reasons)
        self.assertTrue(any("Silent" in reason for reason in reasons))

    def test_report_downloads_pdf(self):
        response = self.client.get(reverse("client_renewal_report"))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/pdf")
        self.assertIn("attachment", response["Content-Disposition"])
        self.assertIn("client_renewal_report_", response["Content-Disposition"])
        self.assertTrue(response["Content-Disposition"].endswith('.pdf"'))
        self.assertTrue(response.content.startswith(b"%PDF-"))
        self.assertGreater(len(response.content), 2000)

    def test_tracker_page_has_report_download(self):
        response = self.client.get(reverse("client_renewal_list"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, reverse("client_renewal_report"))
        self.assertContains(response, "Download report")
        self.assertNotContains(response, "cr-tabs")
        self.assertNotContains(response, "Tracker</a>")


class TicketChartHelperTests(SimpleTestCase):
    def test_format_duration(self):
        self.assertEqual(format_duration(None), "0m")
        self.assertEqual(format_duration(45), "1m")
        self.assertEqual(format_duration(0), "0m")
        self.assertEqual(format_duration(120), "2m")
        self.assertEqual(format_duration(4 * 3600 + 12 * 60), "4h 12m")

    def test_ticket_charts_sla_and_status(self):
        now = timezone.now()
        tickets = [
            empty_ticket_view(
                status="closed",
                created_at=now - timedelta(hours=5),
                completed_at=now - timedelta(hours=1),
                first_responded_at=now - timedelta(hours=4),
                sla={
                    "tracks_resolve": True,
                    "tracks_response": True,
                    "resolve_status": "met",
                    "response_status": "met",
                    "any_breached": False,
                },
            ),
            empty_ticket_view(
                status="open",
                created_at=now - timedelta(hours=10),
                priority="high",
                assignee_id=None,
                sla={"any_breached": True, "tracks_resolve": True, "resolve_status": "breached"},
            ),
        ]
        charts = ticket_charts_from_views(tickets, now=now)
        self.assertEqual(charts["total"], 2)
        self.assertEqual(charts["closed"], 1)
        self.assertEqual(charts["pending"], 1)
        self.assertEqual(charts["sla_breached"], 1)
        self.assertEqual(charts["solved_on_time"], 1)
        self.assertEqual(charts["solved_with_sla"], 1)
        self.assertEqual(charts["solved_on_time_pct"], 100)
        self.assertEqual(charts["queue"]["high_unassigned"], 1)
        self.assertEqual(charts["queue"]["waiting_8h"], 1)
        self.assertEqual(charts["avg_resolution"], "4h")
        self.assertEqual(charts["first_response"], "1h")

    def test_resolution_uses_updated_at_when_completed_missing(self):
        now = timezone.now()
        tickets = [
            empty_ticket_view(
                status="closed",
                created_at=now - timedelta(hours=3),
                completed_at=None,
                updated_at=now - timedelta(hours=1),
            )
        ]
        charts = ticket_charts_from_views(tickets, now=now)
        self.assertEqual(charts["avg_resolution"], "2h")

    def test_resolution_is_zero_minutes_when_none_closed(self):
        now = timezone.now()
        tickets = [
            empty_ticket_view(status="open", created_at=now - timedelta(hours=2))
        ]
        charts = ticket_charts_from_views(tickets, now=now)
        self.assertEqual(charts["avg_resolution"], "0m")
        self.assertEqual(charts["first_response"], "0m")


class SubscriptionImportIdentityTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user(
            "import-owner@example.com", "pass"
        )
        self.company = CompanyProfile.objects.create(
            user=self.user, company_name="Import Co"
        )
        self.client_obj = Client.objects.create(
            company=self.company, name="Acme"
        )
        self.product = Product.objects.create(
            company=self.company, name="Microsoft 365"
        )
        self.vendor = Vendor.objects.create(
            company=self.company, name="Redington"
        )
        self.subscription = ClientSubscription.objects.create(
            company=self.company,
            client=self.client_obj,
            product=self.product,
            vendor=self.vendor,
            subscription_code="SUB-001",
            start_date=date(2026, 1, 1),
            end_date=date(2027, 1, 1),
            quantity=1,
            billing_amount=Decimal("1200"),
        )

    def test_reimport_matches_subscription_id_case_insensitively(self):
        csv_data = (
            "Subscription Id,Product Name,Customer Name,Start Date,"
            "End Date,Quantity,Price\n"
            "sub-001,Microsoft 365,Acme,2026-02-01,2027-02-01,3,2400\n"
        ).encode()

        created, updated, errors = import_subscriptions_from_file(
            self.company,
            csv_data,
            "subscriptions.csv",
            import_vendor=self.vendor,
        )

        self.assertEqual((created, updated, errors), (0, 1, []))
        self.assertEqual(
            ClientSubscription.objects.filter(company=self.company).count(),
            1,
        )
        self.subscription.refresh_from_db()
        self.assertEqual(self.subscription.quantity, 3)
        self.assertEqual(self.subscription.billing_amount, Decimal("2400"))

    def test_import_sets_optional_invoice_number(self):
        csv_data = (
            "Subscription Id,Product Name,Customer Name,Start Date,"
            "End Date,Quantity,Price,Invoice Number\n"
            "SUB-NEW,Microsoft 365,Acme,2026-03-01,2027-03-01,2,1800,INV-42\n"
        ).encode()
        created, updated, errors = import_subscriptions_from_file(
            self.company,
            csv_data,
            "subscriptions.csv",
            import_vendor=self.vendor,
        )
        self.assertEqual((created, updated, errors), (1, 0, []))
        created_sub = ClientSubscription.objects.get(
            company=self.company, subscription_code="SUB-NEW"
        )
        self.assertEqual(created_sub.invoice_number, "INV-42")
        self.assertFalse(created_sub.invoice_pdf)

    def test_database_rejects_case_only_duplicate_subscription_id(self):
        with self.assertRaises(IntegrityError):
            with transaction.atomic():
                ClientSubscription.objects.create(
                    company=self.company,
                    client=self.client_obj,
                    product=self.product,
                    subscription_code="sub-001",
                    start_date=date(2026, 1, 1),
                )


class OemTicketFeatureTests(TestCase):
    def setUp(self):
        User = get_user_model()
        self.user = User.objects.create_user(
            email="oem-owner@example.com", password="pass12345"
        )
        self.company = CompanyProfile.objects.create(
            user=self.user,
            company_name="Acme OEM",
            oem_email="oem-from@example.com",
            oem_email_app_password="oem-secret",
        )
        self.client.login(email="oem-owner@example.com", password="pass12345")
        inbox_patch = patch(
            "business.helpdesk_inbox.maybe_fetch_on_dashboard",
            return_value={"created": 0, "logs": []},
        )
        inbox_patch.start()
        self.addCleanup(inbox_patch.stop)

    def test_company_setup_shows_oem_email(self):
        response = self.client.get(reverse("company_setup"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "OEM email")
        self.assertContains(response, "oem_email")

    def test_internal_and_oem_tabs(self):
        internal = self.client.get(reverse("support_ticket_list"))
        self.assertEqual(internal.status_code, 200)
        self.assertContains(internal, "External (OEM)")
        self.assertContains(internal, "Transfer to OEM")
        oem = self.client.get(reverse("support_ticket_list") + "?source=oem")
        self.assertEqual(oem.status_code, 200)
        self.assertContains(oem, "Email Diary")
        self.assertContains(oem, "Subscription ID")
        self.assertContains(oem, "+ OEM ticket")
        self.assertNotContains(oem, "Serial Number")

    def test_diary_add_company_and_extra_email(self):
        response = self.client.post(
            reverse("support_oem_diary"),
            {"action": "add_company", "name": "HP", "email": "support@hp.example"},
        )
        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertTrue(payload["ok"])
        self.assertEqual(payload["rows"][0]["name"], "HP")
        extra = self.client.post(
            reverse("support_oem_diary"),
            {"action": "add_email", "name": "HP", "email": "care@hp.example"},
        )
        emails = {item["email"] for item in extra.json()["rows"][0]["emails"]}
        self.assertEqual(emails, {"support@hp.example", "care@hp.example"})

    def test_quick_oem_contacts_list_products_with_company_fields(self):
        Product.objects.create(
            company=self.company,
            name="Laptop",
            company_name="HP",
            company_email="support@hp.example",
        )
        Product.objects.create(company=self.company, name="Printer")
        self.client.post(
            reverse("support_oem_diary"),
            {"action": "add_company", "name": "abc", "email": "keniep442006@gmail.com"},
        )
        page = self.client.get(reverse("support_oem_ticket_add") + "?mode=quick")
        self.assertEqual(page.status_code, 200)
        self.assertContains(page, "Search / select product...")
        self.assertContains(page, "Product (optional)")
        response = self.client.get(reverse("support_oem_contacts"))
        self.assertEqual(response.status_code, 200)
        contacts = response.json()["contacts"]
        by_product = {row["product_name"]: row for row in contacts}
        self.assertIn("Laptop", by_product)
        self.assertIn("Printer", by_product)
        self.assertEqual(by_product["Laptop"]["name"], "HP")
        self.assertEqual(by_product["Laptop"]["email"], "support@hp.example")
        self.assertEqual(by_product["Printer"]["name"], "")
        self.assertEqual(by_product["Printer"]["email"], "")
        self.assertNotIn("keniep442006@gmail.com", {row["email"] for row in contacts})

    @patch("business.oem_tickets.send_outbound", return_value=(True, None))
    def test_quick_oem_creates_ticket(self, _send):
        response = self.client.post(
            reverse("support_oem_ticket_add"),
            {
                "mode": "quick",
                "to_email": "vendor@example.com",
                "oem_company": "HP",
                "subject": "Laptop issue",
                "description": "Screen flicker",
                "cc": "",
                "subscription_id": "",
            },
        )
        self.assertEqual(response.status_code, 302)
        from business.models import SupportTicketMeta

        meta = SupportTicketMeta.objects.get(is_oem=True)
        self.assertEqual(meta.oem_company_name, "HP")
        self.assertEqual(meta.oem_to_email, "vendor@example.com")
        self.assertEqual(meta.subscription_id, "")
        html_body = _send.call_args.kwargs["html_body"]
        plain_body = _send.call_args.args[4]
        self.assertIn("OEM SUPPORT", html_body)
        self.assertIn("Laptop issue", html_body)
        self.assertIn("Screen flicker", html_body)
        self.assertIn("Best Regards", html_body)
        self.assertIn("Acme OEM", html_body)
        self.assertNotIn("HP", html_body)
        self.assertIn("Hello,", plain_body)
        self.assertIn("Screen flicker", plain_body)
        self.assertIn("Best Regards,", plain_body)
        self.assertIn("Acme OEM", plain_body)

    def test_oem_email_html_uses_subject_heading_and_attachment_chips(self):
        from types import SimpleNamespace

        from business.oem_tickets import build_oem_email_html, build_oem_email_text

        company = SimpleNamespace(company_name="Acme OEM")
        html_body = build_oem_email_html(
            company,
            subject="Laptop issue",
            body="Screen flicker",
            attachment_names=["photo.png"],
            logo_cid="company-logo",
        )
        self.assertIn("OEM SUPPORT", html_body)
        self.assertIn("<h1", html_body)
        self.assertIn("Laptop issue", html_body)
        self.assertIn("Screen flicker", html_body)
        self.assertNotIn("Attachments", html_body)
        self.assertNotIn("ATTACHMENTS", html_body)
        self.assertIn("photo.png", html_body)
        self.assertIn("cid:company-logo", html_body)
        self.assertIn("Best Regards", html_body)
        self.assertNotIn("ticket id", html_body.lower())
        text_body = build_oem_email_text(
            company,
            subject="Laptop issue",
            body="Screen flicker",
            attachment_names=["photo.png"],
        )
        self.assertIn("- photo.png", text_body)

    @patch("business.oem_tickets.send_outbound", return_value=(True, None))
    def test_transfer_marks_internal_ticket(self, _send):
        from business.helpdesk_bridge import create_internal_ticket

        source = create_internal_ticket(
            company=self.company,
            user=self.user,
            subject="Internal printer issue",
            description="Won't print",
        )
        response = self.client.post(
            reverse("support_ticket_transfer_oem", kwargs={"pk": source.pk}),
            {
                "to_email": "oem@example.com",
                "oem_company": "HP",
                "subject": "Internal printer issue",
                "description": "Won't print",
            },
        )
        self.assertEqual(response.status_code, 302)
        source.crm_meta.refresh_from_db()
        self.assertTrue(source.crm_meta.transferred_to_id)
        detail = self.client.get(
            reverse("support_ticket_detail", kwargs={"pk": source.pk})
        )
        self.assertContains(detail, "Transferred to")

    def test_subscription_lookup_fills_oem_details(self):
        client_obj = Client.objects.create(company=self.company, name="Retail Co")
        product = Product.objects.create(
            company=self.company,
            name="Laptop",
            company_name="HP",
            company_email="support@hp.example",
        )
        ClientSubscription.objects.create(
            company=self.company,
            client=client_obj,
            product=product,
            subscription_code="SUB-88",
            start_date=date(2026, 1, 1),
            serial_number="HP-123",
        )
        response = self.client.get(
            reverse("support_oem_subscription_lookup"), {"q": "SUB-88"}
        )
        self.assertEqual(response.status_code, 200)
        data = response.json()["subscription"]
        self.assertEqual(data["oem_company"], "HP")
        self.assertEqual(data["oem_email"], "support@hp.example")
        self.assertEqual(data["client_name"], "Retail Co")

    def test_subscription_lookup_lists_options_for_dropdown(self):
        client_obj = Client.objects.create(company=self.company, name="Retail Co")
        product = Product.objects.create(
            company=self.company,
            name="Laptop",
            company_name="HP",
            company_email="support@hp.example",
        )
        ClientSubscription.objects.create(
            company=self.company,
            client=client_obj,
            product=product,
            subscription_code="SUB-88",
            start_date=date(2026, 1, 1),
            serial_number="HP-123",
        )
        response = self.client.get(reverse("support_oem_subscription_lookup"))
        self.assertEqual(response.status_code, 200)
        rows = response.json()["subscriptions"]
        self.assertEqual(rows[0]["subscription_id"], "SUB-88")
        self.assertEqual(rows[0]["client_name"], "Retail Co")
        page = self.client.get(reverse("support_oem_ticket_add") + "?mode=standard")
        self.assertEqual(page.status_code, 200)
        self.assertContains(page, "Search / select subscription ID")


class SubscriptionListExportTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user(
            "owner@example.com", "pass"
        )
        self.company = CompanyProfile.objects.create(
            user=self.user, company_name="Co"
        )
        self.client_obj = Client.objects.create(
            company=self.company, name="OMSGuru"
        )
        self.hp_product = Product.objects.create(
            company=self.company,
            name="Laptop",
            category="Hardware",
            company_name="HP",
        )
        self.ms_product = Product.objects.create(
            company=self.company,
            name="Microsoft 365",
            category="Software",
            company_name="Microsoft",
        )
        self.blank_product = Product.objects.create(
            company=self.company,
            name="Unknown Box",
            category="Hardware",
            company_name="",
        )
        self.vendor = Vendor.objects.create(company=self.company, name="Redington")
        self.hp_sub = ClientSubscription.objects.create(
            company=self.company,
            client=self.client_obj,
            product=self.hp_product,
            vendor=self.vendor,
            start_date=date(2026, 1, 1),
            end_date=date(2027, 1, 1),
            quantity=2,
            billing_amount=Decimal("1500.00"),
            serial_number="HP-123",
            subscription_code="SUB-HP",
            invoice_number="INV-1",
            notes="HP laptop note",
        )
        self.ms_sub = ClientSubscription.objects.create(
            company=self.company,
            client=self.client_obj,
            product=self.ms_product,
            vendor=self.vendor,
            start_date=date(2026, 2, 1),
            end_date=date(2027, 2, 1),
            quantity=5,
            billing_amount=Decimal("900.00"),
            subscription_code="SUB-MS",
        )
        self.blank_sub = ClientSubscription.objects.create(
            company=self.company,
            client=self.client_obj,
            product=self.blank_product,
            start_date=date(2026, 3, 1),
            quantity=1,
            billing_amount=Decimal("10.00"),
        )
        self.client.force_login(self.user)

    def test_list_shows_company_column_and_filter(self):
        response = self.client.get(reverse("subscription_list"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "<th>Company</th>", html=True)
        self.assertContains(response, "HP")
        self.assertContains(response, "Microsoft")
        self.assertContains(response, 'name="company_name"')
        self.assertContains(response, "Export")
        self.assertContains(response, "Columns to include")
        self.assertContains(response, "Companies to include")

    def test_filter_by_company_name(self):
        response = self.client.get(
            reverse("subscription_list"), {"company_name": "HP"}
        )
        self.assertEqual(response.status_code, 200)
        names = [row.product.name for row in response.context["object_list"]]
        self.assertEqual(names, ["Laptop"])

    def test_filter_by_blank_company(self):
        response = self.client.get(
            reverse("subscription_list"), {"company_name": NO_COMPANY_VALUE}
        )
        self.assertEqual(response.status_code, 200)
        names = [row.product.name for row in response.context["object_list"]]
        self.assertEqual(names, ["Unknown Box"])

    def test_export_selected_company_and_columns(self):
        response = self.client.post(
            reverse("subscription_export"),
            {
                "columns": ["client", "product", "company", "quantity"],
                "companies": ["HP"],
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        wb = load_workbook(BytesIO(response.content))
        rows = list(wb.active.iter_rows(values_only=True))
        self.assertEqual(rows[0], ("Client", "Product", "Company", "Quantity"))
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[1], ("OMSGuru", "Laptop", "HP", 2))

    def test_export_requires_columns_and_companies(self):
        missing_cols = self.client.post(
            reverse("subscription_export"),
            {"companies": ["HP"]},
        )
        self.assertEqual(missing_cols.status_code, 302)
        missing_companies = self.client.post(
            reverse("subscription_export"),
            {"columns": ["client"]},
        )
        self.assertEqual(missing_companies.status_code, 302)

    def test_export_keeps_page_client_filter(self):
        other_client = Client.objects.create(company=self.company, name="Other Co")
        ClientSubscription.objects.create(
            company=self.company,
            client=other_client,
            product=self.hp_product,
            start_date=date(2026, 4, 1),
            quantity=9,
        )
        response = self.client.post(
            reverse("subscription_export"),
            {
                "client": str(self.client_obj.pk),
                "columns": ["client", "product"],
                "companies": ["HP"],
            },
        )
        self.assertEqual(response.status_code, 200)
        rows = list(
            load_workbook(BytesIO(response.content)).active.iter_rows(values_only=True)
        )
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[1][0], "OMSGuru")

    def test_export_builder_uses_requested_column_order(self):
        self.assertEqual(
            normalize_column_keys(["bogus", "product", "product", "client"]),
            ["product", "client"],
        )
        content = build_subscription_export_bytes(
            [self.ms_sub], ["product", "company"]
        )
        rows = list(load_workbook(BytesIO(content)).active.iter_rows(values_only=True))
        self.assertEqual(rows[0], ("Product", "Company"))
        self.assertEqual(rows[1], ("Microsoft 365", "Microsoft"))


class ProductPriceFieldsTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user(
            "owner@example.com", "pass"
        )
        self.company = CompanyProfile.objects.create(
            user=self.user, company_name="Co"
        )
        self.client.login(email="owner@example.com", password="pass")

    def test_add_form_shows_price_gst_and_discounted_price(self):
        response = self.client.get(reverse("product_add"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Price")
        self.assertContains(response, "GST (%)")
        self.assertContains(response, "Discounted price")
        self.assertContains(response, 'name="unit_price"')
        self.assertContains(response, 'name="tax"')
        self.assertContains(response, 'name="discounted_price"')

    def test_create_product_saves_prices_and_list_shows_unit_price_only(self):
        response = self.client.post(
            reverse("product_add"),
            {
                "name": "Laptop",
                "category": "Hardware",
                "asset_type": "Laptop",
                "unit_price": "1200.50",
                "tax": "18",
                "discounted_price": "999.00",
                "components_count": "0",
            },
            follow=True,
        )
        self.assertEqual(response.status_code, 200)
        product = Product.objects.get(company=self.company, name="Laptop")
        self.assertEqual(product.unit_price, Decimal("1200.50"))
        self.assertEqual(product.tax, "18%")
        self.assertEqual(product.discounted_price, Decimal("999.00"))
        self.assertContains(response, "Price")
        self.assertContains(response, "1200.50")
        self.assertNotContains(response, "Discounted price")
        self.assertNotContains(response, "999.00")


@contextmanager
def _inline_vendor_mail(connection=None):
    """Run the vendor price mail job inline, without a real SMTP login."""
    with patch(
        "business.vendor_quote._start_background", side_effect=lambda fn, *a: fn(*a)
    ), patch("business.vendor_quote.close_old_connections"), patch(
        "business.vendor_quote._uses_graph", return_value=False
    ), patch(
        "business.vendor_quote.open_smtp_connection", return_value=connection
    ) as opener:
        yield opener


class DashboardExpiringTableTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user(
            "owner@example.com", "pass"
        )
        self.company = CompanyProfile.objects.create(
            user=self.user,
            company_name="Co",
            outgoing_email="sales@example.com",
            outgoing_email_app_password="secret",
        )
        self.client_obj = Client.objects.create(company=self.company, name="IMI Tools")
        self.product = Product.objects.create(
            company=self.company, name="Microsoft 365 Apps for business"
        )
        self.vendor = Vendor.objects.create(
            company=self.company, name="Redington", email="vendor@example.com"
        )
        today = timezone.localdate()
        self.subscription = ClientSubscription.objects.create(
            company=self.company,
            client=self.client_obj,
            product=self.product,
            vendor=self.vendor,
            start_date=today - timedelta(days=300),
            end_date=today + timedelta(days=10),
            quantity=2,
            billing_amount=Decimal("100.00"),
            status=ClientSubscription.Status.ACTIVE,
        )
        self.client.force_login(self.user)

    def test_expiring_table_columns_and_win_link(self):
        page = self.client.get(reverse("dashboard"))
        self.assertEqual(page.status_code, 200)
        self.assertContains(page, "Quote to vendor")
        self.assertContains(page, ">Quote<")
        self.assertContains(page, ">Win<")
        self.assertContains(page, ">Lost<")
        self.assertNotContains(page, ">Renew<")
        self.assertContains(page, "required=1")
        renew_url = reverse("subscription_renew", kwargs={"pk": self.subscription.pk})
        self.assertContains(page, renew_url)
        self.assertContains(page, self.product.name)
        html = page.content.decode()
        header = html.split("<thead>", 1)[1].split("</thead>", 1)[0]
        self.assertLess(header.find("End date"), header.find("Product"))
        self.assertLess(header.find("Product"), header.find(">Qty<"))
        self.assertNotIn(">Status<", header)

    def test_quote_popup_maps_product_companies_to_vendors(self):
        Vendor.objects.create(
            company=self.company,
            name="Tata",
            email="tata@example.com",
            catalog_companies=["Reserved VM Instance"],
        )
        self.vendor.catalog_companies = ["microsoft", "Reserved VM Instance"]
        self.vendor.save(update_fields=["catalog_companies"])
        self.product.company_name = "Microsoft"
        self.product.save(update_fields=["company_name"])
        Product.objects.create(
            company=self.company, name="Laptop", company_name="Dell", vendor=self.vendor
        )
        page = self.client.get(reverse("dashboard"))
        self.assertEqual(
            page.context["quote_companies"],
            [
                {"key": "dell", "label": "Dell"},
                {"key": "microsoft", "label": "Microsoft"},
                {"key": "reserved vm instance", "label": "Reserved VM Instance"},
            ],
        )
        by_name = {v["name"]: v["companies"] for v in page.context["quote_vendors"]}
        self.assertEqual(by_name["Redington"], ["microsoft", "reserved vm instance"])
        self.assertEqual(by_name["Tata"], ["reserved vm instance"])
        self.assertContains(page, 'data-company="Microsoft"')
        self.assertContains(page, 'id="dash-company-input"')
        self.assertContains(page, 'id="dash-company-search"')

    def test_lost_from_dashboard_updates_subscription(self):
        response = self.client.post(
            reverse("dashboard_set_deal_lost", kwargs={"pk": self.subscription.pk}),
            {"lost_remark": "price too high"},
        )
        self.assertEqual(response.status_code, 200)
        self.subscription.refresh_from_db()
        self.assertEqual(self.subscription.deal_status, ClientSubscription.DealStatus.LOST)
        self.assertEqual(self.subscription.lost_remark, "price too high")
        page = self.client.get(reverse("dashboard"))
        self.assertNotIn(self.subscription, page.context["expiring_subs"])
        renew_url = reverse("subscription_renew", kwargs={"pk": self.subscription.pk})
        self.assertNotContains(page, renew_url)

    def test_quote_email_uses_price_cta_and_keeps_quantity(self):
        with patch(
            "business.vendor_quote.send_vendor_price_email", return_value=(True, "")
        ) as send, _inline_vendor_mail(), self.captureOnCommitCallbacks(execute=True):
            response = self.client.post(
                reverse("dashboard_quote_vendors", kwargs={"pk": self.subscription.pk}),
                {"vendor_ids": [str(self.vendor.pk)]},
            )
        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertTrue(payload["ok"])
        send.assert_called_once()
        _company, addresses, subject, body, html_body = send.call_args.args[:5]
        self.assertEqual(addresses, ["vendor@example.com"])
        self.assertIn("SUBMIT PRICE", html_body)
        self.assertIn("Microsoft 365 Apps for business", html_body)
        self.assertIn(">Qty</th>", html_body)
        self.assertIn(str(self.subscription.quantity), body)
        end_label = self.subscription.end_date.strftime("%d %b %Y")
        for text in (html_body, body):
            self.assertIn(self.subscription.client.name, text)
            self.assertIn(self.subscription.subscription_list_id, text)
            self.assertIn("Bill Type:", text)
            self.assertIn(end_label, text)
        self.assertIn(
            f"Subscription ID: {self.subscription.subscription_list_id}", body
        )
        self.assertIn(">Subscription ID</th>", html_body)
        self.assertIn(">End Date</th>", html_body)
        self.assertIn(">Days Left</th>", html_body)
        quote = VendorPriceRequest.objects.get(subscription=self.subscription)
        self.assertIsNotNone(quote.email_sent_at)
        public = self.client.get(reverse("vendor_quote_public", kwargs={"token": quote.token}))
        self.assertEqual(public.status_code, 200)
        self.assertContains(public, "Enter price here")
        self.assertContains(public, "Microsoft 365 Apps for business")
        saved = self.client.post(
            reverse("vendor_quote_public", kwargs={"token": quote.token}),
            {"price": "450.50"},
        )
        self.assertEqual(saved.status_code, 200)
        quote.refresh_from_db()
        self.assertEqual(quote.price, Decimal("450.50"))
        self.assertIsNotNone(quote.submitted_at)

    def test_quote_request_returns_before_email_is_sent(self):
        with patch(
            "business.vendor_quote.send_vendor_price_email", return_value=(True, "")
        ) as send:
            response = self.client.post(
                reverse("dashboard_quote_vendors", kwargs={"pk": self.subscription.pk}),
                {"vendor_ids": [str(self.vendor.pk)]},
            )
        self.assertEqual(response.status_code, 200)
        self.assertIn("is being sent to Redington", response.json()["message"])
        send.assert_not_called()
        quote = VendorPriceRequest.objects.get(subscription=self.subscription)
        self.assertIsNone(quote.email_sent_at)

    def test_quote_emails_share_one_smtp_login_and_alert_owner_on_failure(self):
        tata = Vendor.objects.create(
            company=self.company, name="Tata", email="tata@example.com"
        )
        connection = MagicMock()
        results = {"vendor@example.com": (True, ""), "tata@example.com": (False, "550 rejected")}
        with patch(
            "business.vendor_quote.send_vendor_price_email",
            side_effect=lambda company, addresses, *a, **kw: results[addresses[0]],
        ) as send, _inline_vendor_mail(connection) as opener, self.captureOnCommitCallbacks(
            execute=True
        ):
            response = self.client.post(
                reverse("dashboard_quote_vendors", kwargs={"pk": self.subscription.pk}),
                {"vendor_ids": [str(self.vendor.pk), str(tata.pk)]},
            )
        self.assertEqual(response.status_code, 200)
        opener.assert_called_once()
        self.assertEqual(send.call_count, 2)
        self.assertIs(send.call_args_list[0].kwargs["connection"], connection)
        connection.close.assert_called()
        sent = VendorPriceRequest.objects.get(subscription=self.subscription)
        self.assertEqual(sent.vendor, self.vendor)
        self.assertIsNotNone(sent.email_sent_at)
        note = TeamNotification.objects.get(
            recipient=self.user, kind=TeamNotification.Kind.EMAIL_FAILED
        )
        self.assertIn("Tata: 550 rejected", note.body)

    def test_quote_without_smtp_fails_right_away(self):
        self.company.outgoing_email_app_password = ""
        self.company.save(update_fields=["outgoing_email_app_password"])
        response = self.client.post(
            reverse("dashboard_quote_vendors", kwargs={"pk": self.subscription.pk}),
            {"vendor_ids": [str(self.vendor.pk)]},
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn("Configure outgoing SMTP", response.json()["error"])
        self.assertFalse(VendorPriceRequest.objects.exists())

    def test_completed_win_moves_from_dashboard_to_client_renewal(self):
        self.subscription.deal_status = ClientSubscription.DealStatus.WIN
        self.subscription.win_form_completed = True
        self.subscription.save(update_fields=["deal_status", "win_form_completed"])

        dashboard = self.client.get(reverse("dashboard"))
        html = dashboard.content.decode()
        expiring = html.split("Expiring subscriptions", 1)[1].split("Quick actions", 1)[0]
        self.assertNotIn("IMI Tools", expiring)

        renewal = self.client.get(reverse("client_renewal_list"))
        self.assertContains(renewal, "IMI Tools")
        self.assertContains(renewal, self.product.name)


class DashboardSubscriptionRenewalInfoTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user("owner@example.com", "pass")
        self.company = CompanyProfile.objects.create(
            user=self.user,
            company_name="Test Corp",
            outgoing_email="sales@example.com",
            outgoing_email_app_password="secret",
        )
        self.client_obj = Client.objects.create(company=self.company, name="Acme Inc")
        self.product = Product.objects.create(company=self.company, name="Widget Pro")
        self.vendor = Vendor.objects.create(company=self.company, name="Vendor A", email="vendor@example.com")
        today = timezone.localdate()
        self.sub = ClientSubscription.objects.create(
            company=self.company,
            client=self.client_obj,
            product=self.product,
            vendor=self.vendor,
            start_date=today - timedelta(days=300),
            end_date=today + timedelta(days=10),
            quantity=5,
            billing_amount=Decimal("100.00"),
            status=ClientSubscription.Status.ACTIVE,
        )
        self.client.force_login(self.user)

    def test_not_sent(self):
        response = self.client.get(
            reverse("dashboard_subscription_renewal_info", kwargs={"pk": self.sub.pk}),
            HTTP_X_REQUESTED_WITH="XMLHttpRequest",
        )
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertTrue(data["ok"])
        self.assertFalse(data["sent"])
        self.assertIsNone(data["row"])
        self.assertEqual(data["new_products"], [])
        self.assertEqual(data["subscription"]["product_name"], "Widget Pro")
        self.assertEqual(data["subscription"]["id"], self.sub.pk)
        self.assertEqual(data["subscription"]["quantity"], 5)

    def test_sent_not_submitted(self):
        now = timezone.now()
        invite = SubscriptionRenewalInvite.objects.create(
            company=self.company,
            client_email="client@example.com",
            email_sent_at=now,
        )
        SubscriptionRenewalInviteLine.objects.create(
            invite=invite,
            subscription=self.sub,
            company_name="Acme Inc",
            product_name="Widget Pro",
            current_quantity=5,
            renewal_count=None,
        )
        response = self.client.get(
            reverse("dashboard_subscription_renewal_info", kwargs={"pk": self.sub.pk}),
            HTTP_X_REQUESTED_WITH="XMLHttpRequest",
        )
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertTrue(data["ok"])
        self.assertTrue(data["sent"])
        self.assertIsNotNone(data["row"])
        self.assertEqual(data["row"]["subscription_id"], self.sub.pk)
        self.assertIsNone(data["row"]["renewal_quantity"])
        self.assertTrue(len(data["row"]["timestamps"]["sent"]) > 0)
        self.assertEqual(data["new_products"], [])

    def test_sent_and_submitted_with_new_product(self):
        now = timezone.now()
        invite = SubscriptionRenewalInvite.objects.create(
            company=self.company,
            client_email="client@example.com",
            email_sent_at=now,
            submitted_at=now,
            submitted_by_email="client@example.com",
        )
        SubscriptionRenewalInviteLine.objects.create(
            invite=invite,
            subscription=self.sub,
            company_name="Acme Inc",
            product_name="Widget Pro",
            current_quantity=5,
            renewal_count=7,
        )
        SubscriptionRenewalInviteNewProduct.objects.create(
            invite=invite,
            company_name="Acme Inc",
            product_name="Extra Addon",
            quantity=3,
        )
        response = self.client.get(
            reverse("dashboard_subscription_renewal_info", kwargs={"pk": self.sub.pk}),
            HTTP_X_REQUESTED_WITH="XMLHttpRequest",
        )
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertTrue(data["ok"])
        self.assertTrue(data["sent"])
        self.assertEqual(data["row"]["renewal_quantity"], 7)
        self.assertEqual(len(data["new_products"]), 1)
        np = data["new_products"][0]
        self.assertTrue(np["is_new"])
        self.assertEqual(np["product_name"], "Extra Addon")
        self.assertEqual(np["renewal_quantity"], 3)

    def test_only_this_subscription(self):
        product2 = Product.objects.create(company=self.company, name="Other Product")
        sub2 = ClientSubscription.objects.create(
            company=self.company,
            client=self.client_obj,
            product=product2,
            vendor=self.vendor,
            start_date=timezone.localdate() - timedelta(days=300),
            end_date=timezone.localdate() + timedelta(days=10),
            quantity=2,
            status=ClientSubscription.Status.ACTIVE,
        )
        now = timezone.now()
        invite = SubscriptionRenewalInvite.objects.create(
            company=self.company,
            client_email="client@example.com",
            email_sent_at=now,
        )
        SubscriptionRenewalInviteLine.objects.create(
            invite=invite,
            subscription=self.sub,
            company_name="Acme Inc",
            product_name="Widget Pro",
            current_quantity=5,
            renewal_count=5,
        )
        SubscriptionRenewalInviteLine.objects.create(
            invite=invite,
            subscription=sub2,
            company_name="Acme Inc",
            product_name="Other Product",
            current_quantity=2,
            renewal_count=2,
        )
        response = self.client.get(
            reverse("dashboard_subscription_renewal_info", kwargs={"pk": self.sub.pk}),
            HTTP_X_REQUESTED_WITH="XMLHttpRequest",
        )
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data["row"]["subscription_id"], self.sub.pk)
        self.assertEqual(data["row"]["product_name"], "Widget Pro")

    def test_company_isolation(self):
        user2 = get_user_model().objects.create_user("other@example.com", "pass")
        company2 = CompanyProfile.objects.create(
            user=user2,
            company_name="Other Co",
        )
        client2 = Client.objects.create(company=company2, name="Other Client")
        product_other = Product.objects.create(company=company2, name="Other Prod")
        other_sub = ClientSubscription.objects.create(
            company=company2,
            client=client2,
            product=product_other,
            start_date=timezone.localdate(),
            end_date=timezone.localdate() + timedelta(days=5),
            quantity=1,
        )
        response = self.client.get(
            reverse("dashboard_subscription_renewal_info", kwargs={"pk": other_sub.pk}),
            HTTP_X_REQUESTED_WITH="XMLHttpRequest",
        )
        self.assertEqual(response.status_code, 404)

    def test_won_subscriptions_not_included(self):
        ClientSubscription.objects.create(
            company=self.company,
            client=self.client_obj,
            product=self.product,
            start_date=timezone.localdate() - timedelta(days=20),
            end_date=timezone.localdate() + timedelta(days=340),
            quantity=10,
            deal_status=ClientSubscription.DealStatus.WIN,
            deal_status_set_at=timezone.now(),
        )
        now = timezone.now()
        invite = SubscriptionRenewalInvite.objects.create(
            company=self.company,
            client_email="client@example.com",
            email_sent_at=now,
        )
        SubscriptionRenewalInviteLine.objects.create(
            invite=invite,
            subscription=self.sub,
            company_name="Acme Inc",
            product_name="Widget Pro",
            current_quantity=5,
            renewal_count=5,
        )
        response = self.client.get(
            reverse("dashboard_subscription_renewal_info", kwargs={"pk": self.sub.pk}),
            HTTP_X_REQUESTED_WITH="XMLHttpRequest",
        )
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data["row"]["subscription_id"], self.sub.pk)

    def test_client_renewal_list_and_live_regression(self):
        res_list = self.client.get(reverse("client_renewal_list"))
        self.assertEqual(res_list.status_code, 200)
        res_live = self.client.get(reverse("client_renewal_live"), HTTP_X_REQUESTED_WITH="XMLHttpRequest")
        self.assertEqual(res_live.status_code, 200)

    def test_dashboard_renders_plus_button(self):
        page = self.client.get(reverse("dashboard"))
        self.assertEqual(page.status_code, 200)
        info_url = reverse("dashboard_subscription_renewal_info", kwargs={"pk": self.sub.pk})
        self.assertContains(page, f'data-info-url="{info_url}"')
        self.assertContains(page, 'class="dash-info-btn"')
        self.assertContains(page, f'data-client="{self.client_obj.name}"')
        self.assertContains(page, f'data-product="{self.product.name}"')
        self.assertContains(page, 'class="dash-client-cell"')
        self.assertContains(page, ">+<")

    def test_subscription_quote_creates_procurement_deal_and_syncs_prices(self):
        with patch(
            "business.vendor_quote.send_vendor_price_email", return_value=(True, "")
        ) as send, _inline_vendor_mail(), self.captureOnCommitCallbacks(execute=True):
            response = self.client.post(
                reverse("dashboard_quote_vendors", kwargs={"pk": self.sub.pk}),
                {"vendor_ids": [str(self.vendor.pk)]},
            )
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertTrue(data["ok"])
        deal_no = data.get("deal_no")
        self.assertTrue(deal_no.startswith("SQ-"))

        # Verify email contents: no expiry message, no "Expires today", no "Days left"
        send.assert_called_once()
        _company, addresses, subject, body, html_body = send.call_args.args[:5]
        self.assertNotIn("Your subscription is expiring", html_body)
        self.assertNotIn("Expires today", html_body)
        self.assertNotIn("Days left", html_body)
        self.assertIn("SUBMIT PRICE", html_body)
        self.assertIn(deal_no, html_body)

        # Verify ProcurementDeal created
        from business.models import DealVendor, ProcurementDeal
        deal = ProcurementDeal.objects.get(deal_no=deal_no)
        self.assertEqual(deal.product, self.product)
        self.assertEqual(deal.quantity, self.sub.quantity)

        # Vendor submits price via public quote page
        quote = VendorPriceRequest.objects.get(subscription=self.sub)
        saved = self.client.post(
            reverse("vendor_quote_public", kwargs={"token": quote.token}),
            {"price": "750.00"},
        )
        self.assertEqual(saved.status_code, 200)

        # Check procurement price sync
        dv = DealVendor.objects.get(deal_no=deal_no, vendor=self.vendor)
        self.assertEqual(dv.price, Decimal("750.00"))

        # Check dashboard renewal info includes vendor quote
        info_resp = self.client.get(
            reverse("dashboard_subscription_renewal_info", kwargs={"pk": self.sub.pk}),
            HTTP_X_REQUESTED_WITH="XMLHttpRequest",
        )
        self.assertEqual(info_resp.status_code, 200)
        info_data = info_resp.json()
        self.assertEqual(info_data["procurement_deal_no"], deal_no)
        self.assertEqual(len(info_data["vendor_quotes"]), 1)
        self.assertEqual(info_data["vendor_quotes"][0]["vendor_name"], self.vendor.name)
        self.assertEqual(info_data["vendor_quotes"][0]["price"], 750.0)
        self.assertEqual(info_data["vendor_quotes"][0]["status"], "Submitted")

        # A second email to the same vendor shows as its own row, newest first
        with patch(
            "business.vendor_quote.send_vendor_price_email", return_value=(True, "")
        ), _inline_vendor_mail(), self.captureOnCommitCallbacks(execute=True):
            self.client.post(
                reverse("dashboard_quote_vendors", kwargs={"pk": self.sub.pk}),
                {"vendor_ids": [str(self.vendor.pk)]},
            )
        info_data = self.client.get(
            reverse("dashboard_subscription_renewal_info", kwargs={"pk": self.sub.pk}),
            HTTP_X_REQUESTED_WITH="XMLHttpRequest",
        ).json()
        rows = info_data["vendor_quotes"]
        self.assertEqual([r["email_no"] for r in rows], [2, 1])
        self.assertEqual([r["status"] for r in rows], ["Sent", "Submitted"])
        self.assertEqual(rows[1]["price"], 750.0)
        self.assertTrue(rows[0]["sent_at"])


class ClientPanNumberTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user("owner@example.com", "pass")
        self.company = CompanyProfile.objects.create(user=self.user, company_name="Co")
        self.client.force_login(self.user)

    def test_add_client_saves_uppercase_pan_and_lists_it(self):
        response = self.client.post(
            reverse("client_add"),
            {
                "name": "IMI Tools",
                "pan_number": " abcde1234f ",
                "tenant_id": " imitools.onmicrosoft.com ",
                "address": "12 MG Road\nPune",
                "renewal_emails_enabled": "on",
            },
        )
        self.assertEqual(response.status_code, 302)
        client = Client.objects.get(company=self.company, name="IMI Tools")
        self.assertEqual(client.pan_number, "ABCDE1234F")
        self.assertEqual(client.tenant_id, "imitools.onmicrosoft.com")
        self.assertEqual(client.address, "12 MG Road\nPune")
        page = self.client.get(reverse("client_list"))
        self.assertContains(page, "<th>PAN number</th>", html=True)
        self.assertContains(page, "<th>Tenant ID</th>", html=True)
        self.assertContains(page, "<th>Address</th>", html=True)
        self.assertContains(page, "ABCDE1234F")
        self.assertContains(page, "imitools.onmicrosoft.com")
        self.assertContains(page, "12 MG Road<br>Pune")

    def test_add_client_rejects_invalid_pan(self):
        response = self.client.post(
            reverse("client_add"), {"name": "IMI Tools", "pan_number": "12345"}
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Enter a valid PAN (e.g. ABCDE1234F).")
        self.assertFalse(Client.objects.filter(company=self.company).exists())

    def test_excel_export_and_import_handle_pan(self):
        from business.client_import import build_client_export_bytes, import_clients_from_file

        client = Client.objects.create(
            company=self.company,
            name="IMI Tools",
            pan_number="ABCDE1234F",
            tenant_id="tenant-1",
            address="Old address",
        )
        ws = load_workbook(BytesIO(build_client_export_bytes(self.company))).active
        headers = [cell.value for cell in ws[1]]
        self.assertEqual(ws.cell(row=2, column=headers.index("PAN Number") + 1).value, "ABCDE1234F")
        self.assertEqual(ws.cell(row=2, column=headers.index("Tenant ID") + 1).value, "tenant-1")
        self.assertEqual(ws.cell(row=2, column=headers.index("Address") + 1).value, "Old address")

        created, updated, errs = import_clients_from_file(
            self.company,
            b"Name,PAN Number,Tenant ID,Address\n"
            b"IMI Tools,pqrst6789z,tenant-2,New address\nBad Co,XYZ,,\n",
            "clients.csv",
        )
        self.assertEqual((created, updated), (0, 1))
        self.assertEqual(errs, [(3, "Invalid PAN number: XYZ")])
        client.refresh_from_db()
        self.assertEqual(client.pan_number, "PQRST6789Z")
        self.assertEqual(client.tenant_id, "tenant-2")
        self.assertEqual(client.address, "New address")

        import_clients_from_file(self.company, b"Name,GST Number\nIMI Tools,27AAA\n", "c.csv")
        client.refresh_from_db()
        self.assertEqual(client.pan_number, "PQRST6789Z")
        self.assertEqual(client.tenant_id, "tenant-2")
        self.assertEqual(client.address, "New address")
