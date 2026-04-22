"""Parse Excel files for bulk ClientSubscription import."""

from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from accounts.models import CompanyProfile
from business.models import ClientSubscription


HEADER_ALIASES: dict[str, str] = {
    "client": "client",
    "client_name": "client",
    "customer": "client",
    "customer_name": "client",
    "product": "product",
    "product_name": "product",
    "vendor": "vendor",
    "vendor_name": "vendor",
    "quantity": "quantity",
    "qty": "quantity",
    "start_date": "start_date",
    "start": "start_date",
    "end_date": "end_date",
    "end": "end_date",
    "billing_amount": "billing_amount",
    "price": "billing_amount",
    "amount": "billing_amount",
    "invoice_number": "invoice_number",
    "invoice": "invoice_number",
    "serial_number": "serial_number",
    "serial": "serial_number",
    "subscription_code": "subscription_code",
    "subscription_id": "subscription_code",
    "subscriptionid": "subscription_code",
    "status": "status",
    "notes": "notes",
}

ALLOWED_FIELDS = frozenset(HEADER_ALIASES.values())


def _norm_header(val: Any) -> str:
    if val is None:
        return ""
    s = str(val).strip().lower().replace(" ", "_")
    return s


def _cell_str(val: Any) -> str:
    if val is None:
        return ""
    if isinstance(val, float) and val == int(val):
        return str(int(val))
    return str(val).strip()


def _parse_date(val: Any) -> date | None:
    if val is None or val == "":
        return None
    if isinstance(val, datetime):
        return val.date()
    if isinstance(val, date):
        return val
    s = _cell_str(val)
    if not s:
        return None
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def _parse_decimal(val: Any) -> Decimal:
    if val is None or val == "":
        return Decimal("0")
    if isinstance(val, (int, float)):
        return Decimal(str(val))
    s = _cell_str(val).replace(",", "")
    if not s:
        return Decimal("0")
    try:
        return Decimal(s)
    except InvalidOperation:
        return Decimal("0")


def _parse_quantity(val: Any) -> int:
    if val is None or val == "":
        return 1
    if isinstance(val, float) and val == int(val):
        return max(1, int(val))
    if isinstance(val, int):
        return max(1, val)
    s = _cell_str(val)
    if not s:
        return 1
    try:
        return max(1, int(Decimal(s)))
    except (InvalidOperation, ValueError):
        return 1


def _parse_status(raw: Any) -> str:
    s = _cell_str(raw).lower()
    if not s:
        return ClientSubscription.Status.ACTIVE
    by_value = {v: v for v, _ in ClientSubscription.Status.choices}
    if s in by_value:
        return by_value[s]
    by_label = {label.lower(): v for v, label in ClientSubscription.Status.choices}
    return by_label.get(s, ClientSubscription.Status.ACTIVE)


def _build_column_map(header_row: list[Any]) -> dict[str, int]:
    col_by_field: dict[str, int] = {}
    for idx, cell in enumerate(header_row):
        key = _norm_header(cell)
        if not key:
            continue
        field = HEADER_ALIASES.get(key, key if key in ALLOWED_FIELDS else None)
        if field and field not in col_by_field:
            col_by_field[field] = idx
    return col_by_field


def import_subscriptions_from_xlsx(
    company: CompanyProfile,
    file_bytes: bytes,
) -> tuple[int, list[tuple[int, str]]]:
    """
    Returns (created_count, list of (excel_row_number, error_message)).
    """
    errors: list[tuple[int, str]] = []
    created = 0

    if not file_bytes:
        return 0, [(0, "Empty file.")]

    try:
        wb = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        return 0, [(0, f"Could not read Excel file: {exc}")]

    try:
        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)
        try:
            header_row = next(rows_iter)
        except StopIteration:
            return 0, [(0, "The spreadsheet has no rows.")]

        col_map = _build_column_map(list(header_row))
        if "client" not in col_map or "product" not in col_map:
            return 0, [
                (
                    0,
                    "Missing required columns: need headers for client and product "
                    "(e.g. 'client', 'product').",
                )
            ]

        clients_by_lower = {
            c.name.strip().lower(): c for c in company.clients.all()
        }
        products_by_lower = {
            p.name.strip().lower(): p for p in company.products.all()
        }
        vendors_by_lower = {
            v.name.strip().lower(): v for v in company.vendors.all()
        }

        excel_row = 1
        for data_row in rows_iter:
            excel_row += 1
            if data_row is None:
                continue

            def get(field: str) -> Any:
                idx = col_map.get(field)
                if idx is None or idx >= len(data_row):
                    return None
                return data_row[idx]

            client_name = _cell_str(get("client"))
            product_name = _cell_str(get("product"))
            if not client_name and not product_name:
                continue
            if not client_name:
                errors.append((excel_row, "Client is required."))
                continue
            if not product_name:
                errors.append((excel_row, "Product is required."))
                continue

            client = clients_by_lower.get(client_name.lower())
            if not client:
                errors.append(
                    (excel_row, f'Unknown client "{client_name}". Add the client first.')
                )
                continue

            product = products_by_lower.get(product_name.lower())
            if not product:
                errors.append(
                    (
                        excel_row,
                        f'Unknown product "{product_name}". Add the product first.',
                    )
                )
                continue

            vendor_name = _cell_str(get("vendor"))
            vendor = None
            if vendor_name:
                vendor = vendors_by_lower.get(vendor_name.lower())
                if not vendor:
                    errors.append(
                        (
                            excel_row,
                            f'Unknown vendor "{vendor_name}". Leave blank or add the vendor.',
                        )
                    )
                    continue

            start_d = _parse_date(get("start_date"))
            if not start_d:
                errors.append(
                    (excel_row, "Start date is missing or invalid (use YYYY-MM-DD).")
                )
                continue

            end_d = _parse_date(get("end_date"))

            sub = ClientSubscription(
                company=company,
                client=client,
                product=product,
                vendor=vendor,
                status=_parse_status(get("status")),
                start_date=start_d,
                end_date=end_d,
                quantity=_parse_quantity(get("quantity")),
                billing_amount=_parse_decimal(get("billing_amount")),
                invoice_number=_cell_str(get("invoice_number"))[:128],
                serial_number=_cell_str(get("serial_number"))[:128],
                subscription_code=_cell_str(get("subscription_code"))[:64],
                notes=_cell_str(get("notes")),
            )
            sub.save()
            created += 1
    finally:
        wb.close()

    return created, errors
