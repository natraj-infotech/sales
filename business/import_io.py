"""Shared helpers for reading spreadsheet files and detecting header rows."""

from __future__ import annotations

import csv
import logging
import warnings
from collections.abc import Iterator
from io import BytesIO, StringIO
from collections.abc import Callable
from typing import Any

from openpyxl import load_workbook

logger = logging.getLogger(__name__)


def load_workbook_bytes(file_bytes: bytes, *, read_only: bool = False):
    """Load xlsx bytes; suppress harmless openpyxl style warnings."""
    try:
        with warnings.catch_warnings():
            warnings.filterwarnings(
                "ignore",
                message="Workbook contains no default style*",
                category=UserWarning,
            )
            return load_workbook(
                BytesIO(file_bytes),
                read_only=read_only,
                data_only=True,
            )
    except Exception:
        logger.warning("Could not load Excel workbook bytes", exc_info=True)
        raise


def _decode_csv_text(file_bytes: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return file_bytes.decode(encoding)
        except UnicodeDecodeError:
            continue
    return file_bytes.decode("utf-8", errors="replace")


def _read_csv_header_row(file_bytes: bytes) -> list[Any]:
    text = _decode_csv_text(file_bytes)
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(StringIO(text), dialect)
    rows: list[list[str]] = []
    for i, row in enumerate(reader):
        if row:
            rows.append(row)
        if i >= 15:
            break
    if not rows:
        return []
    width = max(len(r) for r in rows)
    header = list(rows[0])
    if len(header) < width:
        header.extend([""] * (width - len(header)))
    return header


def _read_xlsx_header_row(file_bytes: bytes) -> list[Any]:
    """Read every column in row 1 (including empty trailing headers)."""
    wb = load_workbook_bytes(file_bytes, read_only=False)
    try:
        ws = wb.active
        if not ws.max_row or not ws.max_column:
            return []
        return [
            ws.cell(row=1, column=col).value for col in range(1, ws.max_column + 1)
        ]
    finally:
        wb.close()


def read_raw_header_row(file_bytes: bytes, filename: str = "") -> list[Any]:
    lower = (filename or "").lower()
    if lower.endswith(".csv"):
        return _read_csv_header_row(file_bytes)
    if lower.endswith(".xlsx"):
        return _read_xlsx_header_row(file_bytes)
    return []


def cell_str(val: Any) -> str:
    if val is None:
        return ""
    if isinstance(val, float) and val == int(val):
        return str(int(val))
    return str(val).strip()


def read_header_labels(file_bytes: bytes, filename: str = "") -> list[str]:
    """Human-readable label per column index for mapping dropdowns."""
    return [
        cell_str(cell) if cell_str(cell) else f"Column {idx + 1}"
        for idx, cell in enumerate(read_raw_header_row(file_bytes, filename))
    ]


def norm_header(val: Any) -> str:
    if val is None:
        return ""
    return str(val).strip().lower().replace(" ", "_")


def find_best_header_row_index(
    rows: list[tuple[Any, ...] | list[str]],
    header_aliases: dict[str, str],
    hint_words: frozenset[str],
    *,
    norm_fn: Callable[[Any], str] | None = None,
) -> int:
    """Pick the row that looks most like column headers (not a title row)."""
    normalize = norm_fn or norm_header
    best_i = 0
    best_score = -1.0
    for i, row in enumerate(rows[:12]):
        if not row:
            continue
        score = 0.0
        for cell in row:
            key = normalize(cell)
            if not key:
                continue
            if key in header_aliases:
                score += 2.0
            elif any(hint in key for hint in hint_words):
                score += 1.0
            else:
                score += 0.15
        if score > best_score:
            best_score = score
            best_i = i
    return best_i


def detect_header_row(
    file_bytes: bytes,
    filename: str,
    header_aliases: dict[str, str],
    hint_words: frozenset[str],
    *,
    norm_fn: Callable[[Any], str] | None = None,
) -> tuple[int, list[Any]]:
    """Return (header_row_index, padded header row)."""
    rows = peek_spreadsheet_rows(file_bytes, filename)
    if not rows:
        return 0, []
    idx = find_best_header_row_index(
        rows, header_aliases, hint_words, norm_fn=norm_fn
    )
    row = list(rows[idx])
    width = max(len(r) for r in rows)
    if len(row) < width:
        row.extend([None] * (width - len(row)))
    return idx, row


def peek_spreadsheet_rows(
    file_bytes: bytes, filename: str, *, limit: int = 15
) -> list[tuple[Any, ...] | list[str]]:
    return list(_take(iter_spreadsheet_rows(file_bytes, filename), limit))


def _take(iterator: Iterator[Any], limit: int) -> Iterator[Any]:
    for i, item in enumerate(iterator):
        if i >= limit:
            break
        yield item


def iter_spreadsheet_rows(
    file_bytes: bytes,
    filename: str,
    *,
    min_cols: int = 0,
) -> Iterator[tuple[Any, ...] | list[str]]:
    lower = (filename or "").lower()
    if lower.endswith(".csv"):
        return _iter_csv_rows(file_bytes, min_cols=min_cols)
    return _iter_xlsx_rows(file_bytes, min_cols=min_cols)


def _iter_xlsx_rows(
    file_bytes: bytes, *, min_cols: int = 0
) -> Iterator[tuple[Any, ...]]:
    wb = load_workbook_bytes(file_bytes, read_only=False)
    try:
        ws = wb.active
        max_col = max(ws.max_column or 0, min_cols, 1)
        for row in ws.iter_rows(
            min_col=1,
            max_col=max_col,
            values_only=True,
        ):
            yield tuple(row) if row is not None else tuple()
    finally:
        wb.close()


def _iter_csv_rows(
    file_bytes: bytes, *, min_cols: int = 0
) -> Iterator[list[str]]:
    text = _decode_csv_text(file_bytes)
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(StringIO(text), dialect)
    for row in reader:
        if not row:
            continue
        if min_cols and len(row) < min_cols:
            row = row + [""] * (min_cols - len(row))
        yield row
